"""Functional workbook checks using LibreOffice and independent hydraulic references."""
import math
import os
from concurrent.futures import ThreadPoolExecutor
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_pressure_workbook import FIRST, LAST, make_workbook, recalculate


def colebrook_bisection(reynolds, relative_roughness):
    """Solve for x=1/sqrt(f), independently of the workbook fixed-point iteration."""
    if reynolds == 0:
        return 0.0
    if reynolds < 2300:
        return 64 / reynolds
    low, high = 0.01, 100.0
    for _ in range(100):
        x = (low + high) / 2
        residual = x + 2 * math.log10(relative_roughness / 3.7 + 2.51 * x / reynolds)
        if residual > 0:
            high = x
        else:
            low = x
    return 1 / ((low + high) / 2) ** 2


NO_BENDS = {"C7": 0, "C8": 0}
VALID = {
    "standard_water": {},
    "laminar_poiseuille": {"Trykkberegning": {"E24": 100, "E25": 1}, "Bend": NO_BENDS},
    "height_cancel": {"Trykkberegning": {"E30": 12, "E31": 12}},
    "downhill_negative": {"Trykkberegning": {"E30": 0, "E31": 100, "E33": 2}},
    "valves_and_components": {"Ventiler": {"B7": 2, "C7": 4.5, "B8": 1, "C8": 12}, "Komponenter": {"B7": 3, "C7": 7}},
    "y_and_tee_custom_k": {"Bend": {"C9": 3, "F9": 0.7, "C10": 2, "F10": 1.2}},
    "custom_zero_k": {"Bend": {"F7": 0, "F8": 0}},
    "zero_flow_static": {"Trykkberegning": {"E25": 0}},
    "metric_od": {"Trykkberegning": {"E6": "Metrisk OD", "E7": "OD110", "E10": 5.3}},
    "iso_stainless": {"Trykkberegning": {"E6": "ISO 1127", "E7": "DN100", "E10": 2, "E13": "1.4404 / 316L – rustfritt", "E14": "EN 10216-5"}},
    "actual_id_override": {"Trykkberegning": {"E11": 75}},
    "own_medium": {"Trykkberegning": {"E20": "Eget medium", "E23": 1100, "E24": 2.5}},
    "flange_full_bore": {"Trykkberegning": {"E17": 10, "E18": 0}},
    "flange_custom_k": {"Trykkberegning": {"E17": 10, "E18": 0.1}},
    "laminar_with_documented_k": {"Trykkberegning": {"E24": 100, "E25": 1}, "Bend": {"F7": 1.2, "F8": 0.6}},
    # Published fluids/Rennels example: Di=4.020 m, R=20 m, angle=30°, Re=1e5.
    "rennels_published_example": {
        "Trykkberegning": {"E9": 5000, "E11": 4020, "E16": 0, "E23": 1000, "E24": 1,
                           "E25": 100000 * 0.001 * math.pi * 4.02 / (4 * 1000) * 3600,
                           "E29": 0, "E30": 0, "E31": 0},
        "Bend": {"C7": 1, "B7": 30, "D7": 20 / 4.02, "C8": 0},
    },
}
DEFAULT_DI = (88.9 - 2 * 3.2) / 1000
INVALID = {
    "negative_flow": {"Trykkberegning": {"E25": -1}},
    "text_flow": {"Trykkberegning": {"E25": "ikke et tall"}},
    "missing_flow": {"Trykkberegning": {"E25": None}},
    "invalid_diameter": {"Trykkberegning": {"E10": 50}},
    "id_larger_than_od": {"Trykkberegning": {"E11": 100}},
    "metric_wrong_dimension": {"Trykkberegning": {"E6": "Metrisk OD", "E7": "DN80"}},
    "material_standard_mismatch": {"Trykkberegning": {"E14": "EN 10216-5"}},
    "own_medium_missing_properties": {"Trykkberegning": {"E20": "Eget medium"}},
    "negative_viscosity": {"Trykkberegning": {"E24": -1}},
    "missing_tee_k": {"Bend": {"C9": 1}},
    "invalid_bend_angle": {"Bend": {"B7": 181}},
    "invalid_bend_radius": {"Bend": {"D7": 0.3}},
    "negative_k": {"Bend": {"F7": -1}},
    "fractional_quantity": {"Bend": {"C7": 1.5}},
    "valve_missing_drop": {"Ventiler": {"B7": 1}},
    "component_negative_drop": {"Komponenter": {"B7": 1, "C7": -2}},
    "flange_missing_k": {"Trykkberegning": {"E17": 5, "E18": None}},
    "negative_flange_count": {"Trykkberegning": {"E17": -1}},
    "negative_height_input": {"Trykkberegning": {"E30": -1}},
    "laminar_automatic_bends": {"Trykkberegning": {"E24": 100, "E25": 1}},
    "transition": {"Trykkberegning": {"E25": 3000 * 0.001002 * math.pi * DEFAULT_DI / (4 * 998.2) * 3600}},
    "zero_flow_with_loss": {"Trykkberegning": {"E25": 0}, "Ventiler": {"B7": 1, "C7": 4}},
}


class PressureWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory(prefix="pressure-regression-")
        cls.directory = Path(cls.temporary.name)
        source = cls.directory / "input"
        source.mkdir()
        paths = []
        all_cases={**VALID,**INVALID}
        selected=os.environ.get("PRESSURE_TEST_CASES")
        cls.case_names=set(selected.split(",")) if selected else set(all_cases)
        if not cls.case_names<=set(all_cases):raise ValueError("Unknown PRESSURE_TEST_CASES")
        cls.case_names.add("standard_water")
        for name, updates in all_cases.items():
            if name not in cls.case_names:continue
            workbook = make_workbook()
            for sheet, values in updates.items():
                for cell, value in values.items():
                    workbook[sheet][cell] = value
            path = source / f"{name}.xlsx"
            workbook.save(path)
            workbook.close()
            paths.append(path)
        # Independent profiles permit bounded parallelism without a shared Calc process.
        workers=min(4,len(paths))
        def calculate_group(index):
            recalculate(paths[index::workers], cls.directory / "calculated", cls.directory / f"profile-{index}", timeout=900)
        with ThreadPoolExecutor(max_workers=workers) as executor:
            list(executor.map(calculate_group, range(workers)))

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def workbook(self, name, formulas=False):
        return load_workbook(self.directory / "calculated" / f"{name}.xlsx", data_only=not formulas)

    def close_enough(self, actual, expected):
        self.assertIsInstance(actual, (int, float))
        self.assertTrue(math.isclose(actual, expected, rel_tol=2e-6, abs_tol=1e-7), (actual, expected))

    def no_errors(self, workbook):
        errors = [(sheet.title, cell.coordinate, cell.value) for sheet in workbook
                  for row in sheet for cell in row if cell.data_type == "e"]
        self.assertEqual(errors, [])

    def test_valid_hydraulic_scenarios(self):
        for name in VALID:
            if name not in self.case_names:continue
            with self.subTest(case=name):
                w = self.workbook(name)
                self.addCleanup(w.close)
                self.no_errors(w)
                main, calc, bends = w["Trykkberegning"], w["Beregning"], w["Bend"]
                self.assertEqual(calc["B4"].value, "OK")
                rho, mu, diameter = main["E21"].value, main["E22"].value / 1000, main["E12"].value / 1000
                flow = main["E25"].value / 3600
                velocity = flow / (math.pi * diameter**2 / 4)
                reynolds = rho * velocity * diameter / mu
                friction = colebrook_bisection(reynolds, main["E15"].value / 1000 / diameter)
                dynamic = rho * velocity**2 / 2
                pipe_loss = friction * main["E29"].value / diameter * dynamic
                if name == "laminar_poiseuille":
                    pipe_loss = 128 * mu * main["E29"].value * flow / (math.pi * diameter**4)
                bend_loss = 0
                for r in range(FIRST, LAST + 1):
                    count = bends[f"C{r}"].value or 0
                    if not count:
                        continue
                    own_k = bends[f"F{r}"].value
                    if own_k is not None:
                        k = own_k
                    else:
                        angle = math.radians(bends[f"B{r}"].value)
                        radius = bends[f"D{r}"].value or (1.5 if bends[f"A{r}"].value == "Bend, lang radius" else 1)
                        sine = math.sin(angle / 2)
                        k = friction * angle * radius + (0.1 + 2.4 * friction) * sine + 6.6 * friction * (math.sqrt(sine) + sine) / radius**(4 * angle / math.pi)
                    self.close_enough(bends[f"G{r}"].value, k)
                    bend_loss += count * k * dynamic
                fixed = []
                for sheet in ("Ventiler", "Komponenter"):
                    fixed.append(sum((w[sheet][f"B{r}"].value or 0) * (w[sheet][f"C{r}"].value or 0) * 1000 for r in range(FIRST, LAST + 1)))
                flange_loss = main["E17"].value * main["E18"].value * dynamic
                static = rho * 9.80665 * (main["E30"].value - main["E31"].value)
                dissipative = pipe_loss + bend_loss + sum(fixed) + flange_loss
                expected = {13: velocity, 14: reynolds, 23: friction, 29: pipe_loss, 30: bend_loss,
                            31: fixed[0], 32: fixed[1], 33: flange_loss, 35: dissipative,
                            36: static, 37: dissipative + static}
                for row, number in expected.items():
                    self.close_enough(calc[f"B{row}"].value, number)
                self.close_enough(main["H7"].value, (dissipative + static) / 100000)
                if name == "downhill_negative":
                    self.assertLess(main["H7"].value, 0)
                    self.close_enough(main["L19"].value, 2 - (dissipative + static) / 100000)
                if name == "metric_od":
                    self.close_enough(main["E8"].value, 110)
                    self.close_enough(main["E12"].value, 99.4)
                if name == "iso_stainless":
                    self.close_enough(main["E8"].value, 114.3)
                    self.close_enough(main["E15"].value, 0.015)
                if name == "rennels_published_example":
                    self.close_enough(bends["G7"].value, 0.11519070808085191)
                print(f"PASS hydraulics: {name}")

    def test_invalid_inputs_block_result(self):
        if not self.case_names.intersection(INVALID):self.skipTest("No invalid scenarios selected for this focused run")
        for name in INVALID:
            if name not in self.case_names:continue
            with self.subTest(case=name):
                w = self.workbook(name)
                self.addCleanup(w.close)
                self.no_errors(w)
                self.assertTrue(w["Beregning"]["B4"].value.startswith("FEIL:"), w["Beregning"]["B4"].value)
                self.assertIn(w["Trykkberegning"]["H7"].value, (None, ""))
                self.assertIn(w["Beregning"]["B37"].value, (None, ""))
                print(f"PASS input guard: {name}")

    def test_lists_formulas_and_structure_survive_recalculation(self):
        w = self.workbook("standard_water", formulas=True)
        self.addCleanup(w.close)
        self.assertEqual(w.active.title, "Trykkberegning")
        self.assertEqual(w["_Lister"].sheet_state, "hidden")
        for name in ("Rorserier", "NominelleDN", "MetriskeOD", "Materialvalg", "Medievalg", "ENStandarder", "Bendtyper"):
            self.assertIn(name, w.defined_names)
        for sheet in ("Trykkberegning", "Bend", "Ventiler", "Komponenter"):
            self.assertTrue(w[sheet].data_validations.dataValidation)
        validations = [dv for dv in w["Trykkberegning"].data_validations.dataValidation if "E7" in dv]
        self.assertEqual(len(validations), 1)
        self.assertIn("INDIRECT", validations[0].formula1)
        self.assertEqual(len(w["Trykkberegning"]._charts), 1)
        self.assertEqual(w["Trykkberegning"]["H7"].data_type, "f")
        for sheet in w:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        self.assertNotIn("#REF!", cell.value)

    def test_delivered_workbook_has_complete_cached_results(self):
        artifact = ROOT / "workbooks" / "Trykkberegning_ror.xlsx"
        self.assertTrue(artifact.is_file())
        w = load_workbook(artifact, data_only=True)
        self.addCleanup(w.close)
        self.no_errors(w)
        self.assertEqual(w["Beregning"]["B4"].value, "OK")
        self.assertGreater(w["Beregning"]["B30"].value, 0)
        self.close_enough(w["Trykkberegning"]["H7"].value, w["Beregning"]["B37"].value / 100000)


if __name__ == "__main__":
    unittest.main(verbosity=2)
