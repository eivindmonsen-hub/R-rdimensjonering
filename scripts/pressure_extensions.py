"""Pressure-only pipe strength, serial reducers, pipe sections and valve catalog UI."""
import csv
from pathlib import Path

from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
from openpyxl.styles import Alignment, PatternFill

if __package__:
    from .build_pressure_workbook import (
        ROOT, FIRST, LAST, NAVY, TEAL, BLUE, PALE, WHITE, INK, GREY, AMBER, RED,
        base_sheet, merge, styled_cell, dv_list, dv_number, named,
    )
else:
    from build_pressure_workbook import (
        ROOT, FIRST, LAST, NAVY, TEAL, BLUE, PALE, WHITE, INK, GREY, AMBER, RED,
        base_sheet, merge, styled_cell, dv_list, dv_number, named,
    )


def condition_chain(checks, success='"OK"'):
    result=success
    for condition,message in reversed(checks):result=f'IF({condition},"{message}",{result})'
    return result


def guarded(expression, guard):
    return f'=IFERROR(IF({guard},{expression},""),"")'


def strength(wb):
    ws=base_sheet(wb,"Rørstyrke","RØRSTYRKE – INDRE OVERTRYKK","Lamé + von Mises · rett sylinder med lukkede ender · materialdata ved dimensjonerende temperatur",[18,12,12,14,28,15,3,17,14,14,14,16,16])
    merge(ws,"A5:F5","INNDATA / DOKUMENTERTE VERDIER",WHITE,TEAL,11,True)
    entries={
        6:("Dimensjonerende overtrykk",10,"bar","Hydraulisk trykktap er ikke dimensjonerende overtrykk. Oppgi maksimal trykkforskjell mellom innside og omgivelser, inkludert relevante trykktopper."),
        7:("Ytre diameter", "='Trykkberegning'!E8","mm","Koblet til hovedfanen; kan overstyres med dokumentert ytre diameter."),
        8:("Nominell / målt godstykkelse","='Trykkberegning'!E10","mm","Bruk metallgodset, ikke hydraulisk ID av et belegg. Egen hydraulisk ID i hovedfanen brukes ikke til styrkegods."),
        9:("Negativ godstoleranse",0,"%","Oppgi dokumentert negativ produksjonstoleranse. 0 kan brukes ved garantert eller målt minste godstykkelse, uten å telle toleransen to ganger."),
        10:("Korrosjons-/erosjonstillegg",0,"mm","Trekkes fra tilgjengelig gods etter negativ toleranse."),
        12:("Dimensjonerende temperatur",20,"°C","Materialets flytegrense må dokumenteres ved denne temperaturen. Ingen automatisk temperaturinterpolasjon."),
        13:("Flytegrense Re / Rp0,2 ved T",None,"MPa","Obligatorisk dokumentert materialverdi ved valgt temperatur. Ingen antatt universell flytegrense for materialnavnet."),
        14:("Materialdata: kilde / sertifikat",None,"","Referanse til materialnorm, temperaturtabell eller materialsertifikat."),
        15:("Krav til sikkerhetsfaktor n",1.5,"–","Bruk prosjekteringskravet mot flyt. 1,5 er en eksempelverdi, ikke et universelt normkrav."),
        16:("Styrkereduksjonsfaktor η",1,"–","Dokumentert faktor 0<η<=1, f.eks. for skjøt. Dette er en enkel styrkereduksjon, ikke full EN-beregning av sveis."),
        17:("Data gjelder valgt temperatur","Nei","","Sett Ja først når materialstyrken gjelder valgt materiale, tykkelse og temperatur."),
    }
    for r,(label,value,unit,note) in entries.items():
        merge(ws,f"A{r}:D{r}",label,size=10)
        styled_cell(ws,f"E{r}",value,"input","0.000",note)
        if unit:styled_cell(ws,f"F{r}",unit)
        else:
            ws.merge_cells(f"E{r}:F{r}")
            ws.row_dimensions[r].height=38
        ws.row_dimensions[r].height=max(ws.row_dimensions[r].height,28)
    merge(ws,"A11:D11","Effektivt tilgjengelig gods",size=10)
    styled_cell(ws,"E11","=IFERROR(E8*(1-E9/100)-E10,0)",number_format="0.000")
    styled_cell(ws,"F11","mm")
    merge(ws,"A19:F20","Materialvalg fra hovedfanen: ",INK,PALE)
    ws["A19"]="=\"Materialvalg fra hovedfanen: \"&'Trykkberegning'!E13"
    checks=[
        ('OR(NOT(ISNUMBER(E6)),E6<0,E6="")',"FEIL: dimensjonerende overtrykk må være >=0."),
        ('OR(NOT(ISNUMBER(E7)),E7<=0,NOT(ISNUMBER(E8)),E8<=0,E8>=E7/2)',"FEIL: kontroller ytre diameter og metallgods."),
        ('OR(NOT(ISNUMBER(E9)),E9<0,E9>=100,NOT(ISNUMBER(E10)),E10<0,E11<=0)',"FEIL: kontroller toleranse og korrosjonstillegg."),
        ('OR(NOT(ISNUMBER(E12)),E12<=-273.15)',"FEIL: oppgi dimensjonerende temperatur."),
        ('OR(NOT(ISNUMBER(E13)),E13<=0,E14="",E17<>"Ja")',"MANGLER: dokumentert flytegrense ved temperatur, kilde og bekreftelse."),
        ('OR(NOT(ISNUMBER(E15)),E15<1,NOT(ISNUMBER(E16)),E16<=0,E16>1)',"FEIL: n må være >=1 og η må være >0 og <=1."),
    ]
    styled_cell(ws,"E39",f'=IFERROR({condition_chain(checks)},"FEIL: kontroller tallfeltene.")')
    merge(ws,"H5:M5","STYRKEKONTROLL / KUN INDRE OVERTRYKK",WHITE,NAVY,11,True)
    merge(ws,"H6:M8","=E39",INK,AMBER,11,True)
    merge(ws,"H10:M10","Modellens trykkgrense ved valgt sikkerhetsfaktor",GREY,size=10)
    c=merge(ws,"H11:M13",'=IF(E39="OK",E30,"")',WHITE,TEAL,27,True)
    c.number_format='0.00" bar"'
    result_rows=[
        (15,"Von Mises-spenning","E26",'0.000" MPa"'),
        (16,"Tillatt spenning = ηRe / n","E27",'0.000" MPa"'),
        (17,"Beregnet faktor mot flyt","E28",'0.000'),
        (18,"Utnyttelse","E29",'0.0%'),
        (19,"Modellens minste effektive gods","E31",'0.000" mm"'),
        (20,"Modellens minste nominelle gods","E32",'0.000" mm"'),
    ]
    for r,label,ref,fmt in result_rows:
        merge(ws,f"H{r}:K{r}",label,size=10)
        c=merge(ws,f"L{r}:M{r}",f'=IF(E39="OK",{ref},"")',INK,PALE,11,True)
        c.number_format=fmt
    merge(ws,"H22:M24",'=IF(E39<>"OK",E39,IF(E29<=1,"KRAV OPPFYLT i idealisert modell for indre overtrykk.","KRAV IKKE OPPFYLT ved valgt sikkerhetsfaktor."))',INK,AMBER,11,True)
    values={
        21:("Ytre radius ro","E7/2","mm"),
        22:("Indre radius ri, effektivt gods","E21-E11","mm"),
        23:("Maks ringspenning, indre overflate","(E6/10)*(E21^2+E22^2)/(E21^2-E22^2)","MPa"),
        24:("Aksialspenning, lukkede ender","(E6/10)*E22^2/(E21^2-E22^2)","MPa"),
        25:("Radialspenning, indre overflate","-E6/10","MPa"),
        26:("Ekvivalent von Mises-spenning","SQRT(((E23-E24)^2+(E24-E25)^2+(E25-E23)^2)/2)","MPa"),
        27:("Tillatt spenning","E13*E16/E15","MPa"),
        28:("Beregnet faktor mot flyt",'IF(E26=0,"∞ (ingen trykklast)",E13*E16/E26)',"–"),
        29:("Utnyttelse","E26/E27","–"),
        30:("Modellens trykkgrense","10*E27*(E21^2-E22^2)/(SQRT(3)*E21^2)","bar"),
        31:("Nødvendig effektivt gods",'IF(SQRT(3)*(E6/10)/E27>=1,"Ingen elastisk løsning",E21*(1-SQRT(1-SQRT(3)*(E6/10)/E27)))',"mm"),
        32:("Nødvendig nominelt gods",'IF(ISNUMBER(E31),(E31+E10)/(1-E9/100),"Ingen elastisk løsning")',"mm"),
    }
    for r,(label,expr,unit) in values.items():
        merge(ws,f"A{r}:D{r}",label,size=10)
        styled_cell(ws,f"E{r}",guarded(expr,'E39="OK"'),number_format="0.000000")
        styled_cell(ws,f"F{r}",unit)
    merge(ws,"A35:M37","OMFANG: lineær-elastisk, homogen rett sylinder med lukkede ender og indre overtrykk. Lamés spenninger ved indre overflate og von Mises mot dokumentert flytegrense. Dette er ikke full verifikasjon etter EN 13480: reduksjoner, flenser, bøying, termiske laster, utmatting, kryp og stabilitet ved utvendig trykk inngår ikke. Dimensjonerende overtrykk oppgis separat fra hydraulisk trykktap.",GREY,PALE,10)
    for r in [6,7,8,9,10,13,15,16]:dv_number(ws,f"E{r}",0 if r not in (7,8,13,15) else 1e-9,100 if r==9 else 1 if r==16 else 1e9)
    dv_number(ws,"E12",-273.14,10000)
    dv_list(ws,'"Nei,Ja"',"E17")
    ws.row_dimensions[39].height=60
    ws.print_area="A1:M39"


def reducers(wb):
    ws=base_sheet(wb,"Reduksjoner","KONSENTRISKE REDUKSJONER / EKSPANSJONER","Registrer i strømningsrekkefølge. Én aktiv komponent per rad; «med» er 0 eller 1.",[30,9,18,18,17,17,16,16,17,19,42,40])
    merge(ws,"A4:B4","Innløpets ID [mm]",size=10)
    styled_cell(ws,"C4","='Trykkberegning'!E12",number_format="0.00")
    merge(ws,"D4:E4","Utløpets ID [mm]",size=10)
    styled_cell(ws,"F4",f'=IFERROR(LOOKUP(2,1/(B{FIRST}:B{LAST}=1),D{FIRST}:D{LAST}),C4)',number_format="0.00")
    merge(ws,"G4:L4","Endret hastighet beregnes én gang fra faktisk innløps-ID og siste aktive utløps-ID.",INK,BLUE,10)
    merge(ws,"A5:L5","Crane-estimat for turbulent strømning, K referert til den minste diameteren. Eget K overstyrer. Rett rørfriksjon ved ulike dimensjoner legges i Rørstrekk.",INK,AMBER,10)
    headers=["Beskrivelse","Med 0/1","ID inn [mm]","ID ut [mm]","Lengde [mm]","Full konusvinkel [°]","K modell","Eget K","Hastighet liten [m/s]","Lokalt tap [kPa]","Kontroll","Datakilde / kommentar"]
    for col,label in enumerate(headers,1):styled_cell(ws,f"{chr(64+col)}6",label)
    ws.row_dimensions[6].height=38
    for r in range(FIRST,LAST+1):
        for col,value in {"A":"Konsentrisk reduksjon" if r==FIRST else None,"B":0,"C":"='Trykkberegning'!E12" if r==FIRST else None,"D":None,"E":None,"H":None,"L":None}.items():
            styled_cell(ws,f"{col}{r}",value,"input","0" if col=="B" else "0.000")
        small,large=f"MIN(C{r},D{r})",f"MAX(C{r},D{r})"
        beta=f"({small}/{large})"
        styled_cell(ws,f"F{r}",f'=IFERROR(IF(B{r}=1,IF(E{r}=0,180,2*DEGREES(ATAN(ABS(C{r}-D{r})/(2*E{r})))),""),"")',number_format="0.000")
        sin=f"SIN(RADIANS(F{r})/2)"
        contraction=f'IF(F{r}<45,0.8*{sin}*(1-{beta}^2),0.5*SQRT({sin})*(1-{beta}^2))'
        expansion=f'IF(F{r}<45,2.6*{sin}*(1-{beta}^2)^2,(1-{beta}^2)^2)'
        styled_cell(ws,f"G{r}",f'=IFERROR(IF(AND(B{r}=1,C{r}>0,D{r}>0,C{r}<>D{r}),IF(C{r}>D{r},{contraction},{expansion}),""),"")',number_format="0.00000")
        styled_cell(ws,f"I{r}",guarded(f'Beregning!$B$7/(PI()*({small}/1000)^2/4)',f'AND(B{r}=1,Beregning!$B$3="OK",C{r}>0,D{r}>0)'),number_format="0.000")
        upstream="'Trykkberegning'!$E$12" if r==FIRST else f'IFERROR(LOOKUP(2,1/($B${FIRST}:B{r-1}=1),$D${FIRST}:D{r-1}),\'Trykkberegning\'!$E$12)'
        local_re=f'Beregning!$B$10*I{r}*({small}/1000)/Beregning!$B$9'
        checks=[
            (f'OR(B{r}="",B{r}=0)',"Ikke med"),
            (f'OR(NOT(ISNUMBER(B{r})),B{r}<>1)',"FEIL: med må være 0 eller 1"),
            (f'OR(NOT(ISNUMBER(C{r})),C{r}<=0,NOT(ISNUMBER(D{r})),D{r}<=0,C{r}=D{r})',"FEIL: kontroller diametre"),
            (f'ABS(C{r}-({upstream}))>0.000001',"FEIL: ID inn må følge forrige utløp"),
            (f'OR(E{r}="",NOT(ISNUMBER(E{r})),E{r}<0)',"FEIL: oppgi lengde >=0"),
            ('Beregning!$B$3<>"OK"',"Venter på inndata"),
            (f'AND(H{r}<>"",OR(NOT(ISNUMBER(H{r})),H{r}<0))',"FEIL: eget K"),
            (f'AND(H{r}="",{local_re}>0,{local_re}<4000)',"FEIL: oppgi K ved lavt Re"),
        ]
        styled_cell(ws,f"K{r}",f'=IFERROR({condition_chain(checks)},"FEIL: kontroller tall")')
        styled_cell(ws,f"J{r}",guarded(f'IF(H{r}<>"",H{r},G{r})*Beregning!$B$10*I{r}^2/2000',f'K{r}="OK"'),number_format="0.000")
    dv_number(ws,f"B{FIRST}:B{LAST}",0,1,True)
    for col in ("C","D","E","H"):dv_number(ws,f"{col}{FIRST}:{col}{LAST}")
    merge(ws,"A33:I33","SUM IRREVERSIBELT LOKALT TAP",WHITE,TEAL,11,True)
    styled_cell(ws,"J33",f"=SUM(J{FIRST}:J{LAST})",number_format='0.000" kPa"')
    merge(ws,"A35:L37","Modellen gjelder en seriell strømningsvei med samme volumstrøm. Diametrene må følge hverandre i registrert rekkefølge. Lokalt K-tap er alltid positivt; hastighetsbidraget ρ(vut²−vinn²)/2 har fortegn og tas med én gang i hovedresultatet. Hoveddimensjon-modus legger alle rette meter før første reduksjon. Ved rette rør mellom/etter reduksjoner bruker du Rørstrekk-tabell. Crane er et geometrisk anslag, ikke et leverandørdatablad; dokumentert K foretrekkes. Lengde 0 representerer en skarp overgang.",GREY,PALE,10)
    ws.print_area="A1:L37"


def pipe_sections(wb):
    ws=base_sheet(wb,"Rørstrekk","RETTE RØRSTREKK MED ULIKE DIAMETRE","Aktiver «Rørstrekk-tabell» i hovedfanen. Da erstatter tabellen hovedfanens rette rørmeter.",[29,18,19,19,17,19,16,22,42,45])
    merge(ws,"A4:J5","Oppgi hver rette lengde med faktisk innvendig diameter. Ruhet kan overstyres per strekk. Alle strekk bruker samme væske og volumstrøm. Bend og reduksjoners buede/koniske lengder skal ikke telles to ganger.",INK,BLUE,10)
    headers=["Strekk","Lengde [m]","ID [mm]","Egen ruhet [mm]","Hastighet [m/s]","Reynolds","Darcy f","Rørtap [kPa]","Kontroll","Kommentar"]
    for col,label in enumerate(headers,1):styled_cell(ws,f"{chr(64+col)}6",label)
    ws.row_dimensions[6].height=36
    for r in range(FIRST,LAST+1):
        for col,value in {"A":"Hovedstrekk" if r==FIRST else None,"B":"='Trykkberegning'!E29" if r==FIRST else 0,"C":"='Trykkberegning'!E12" if r==FIRST else None,"D":None,"J":None}.items():styled_cell(ws,f"{col}{r}",value,"input","0.000")
        eps=f'IF(D{r}="",\'Trykkberegning\'!$E$15,D{r})'
        styled_cell(ws,f"E{r}",guarded(f'Beregning!$B$7/(PI()*(C{r}/1000)^2/4)',f'AND(Beregning!$B$3="OK",C{r}>0)'),number_format="0.000")
        styled_cell(ws,f"F{r}",guarded(f'Beregning!$B$10*E{r}*(C{r}/1000)/Beregning!$B$9',f'AND(Beregning!$B$3="OK",C{r}>0)'),number_format="0")
        re,rough=f"F{r}",f"({eps}/C{r})"
        styled_cell(ws,f"L{r}",f'=IFERROR(IF({re}=0,0,IF({re}<2300,64/{re},(-1.8*LOG10(({rough}/3.7)^1.11+6.9/{re}))^(-2))),"")')
        for col in range(13,19):
            letter,previous=chr(64+col),chr(63+col)
            styled_cell(ws,f"{letter}{r}",f'=IFERROR(IF({re}=0,0,IF({re}<2300,L{r},(-2*LOG10({rough}/3.7+2.51/({re}*SQRT({previous}{r}))))^(-2))),"")')
            ws.column_dimensions[letter].hidden=True
        ws.column_dimensions["L"].hidden=True
        styled_cell(ws,f"G{r}",f"=R{r}",number_format="0.00000")
        checks=[
            (f'OR(B{r}="",B{r}=0)',"Ikke med"),
            (f'OR(NOT(ISNUMBER(B{r})),B{r}<0,NOT(ISNUMBER(C{r})),C{r}<=0)',"FEIL: lengde eller diameter"),
            (f'OR(NOT(ISNUMBER({eps})),{eps}<0,{eps}>=C{r})',"FEIL: ruhet"),
            ('Beregning!$B$3<>"OK"',"Venter på inndata"),
            (f'AND(F{r}>=2300,F{r}<4000)',"FEIL: overgangsområde"),
            (f'AND(ABS(C{r}-\'Trykkberegning\'!$E$12)>0.000001,COUNTIFS(Reduksjoner!$B${FIRST}:$B${LAST},1,Reduksjoner!$D${FIRST}:$D${LAST},C{r})=0)',"FEIL: diameter mangler i reduksjonskjeden"),
        ]
        styled_cell(ws,f"I{r}",f'=IFERROR({condition_chain(checks)},"FEIL: kontroller tall")')
        styled_cell(ws,f"H{r}",guarded(f'G{r}*B{r}/(C{r}/1000)*Beregning!$B$10*E{r}^2/2000',f'I{r}="OK"'),number_format="0.000")
    for col in ("B","C","D"):dv_number(ws,f"{col}{FIRST}:{col}{LAST}")
    merge(ws,"A33:G33","SUM RØRFRIKSJON – ERSTATTER HOVEDLENGDEN NÅR AKTIVERT",WHITE,TEAL,11,True)
    styled_cell(ws,"H33",f"=SUM(H{FIRST}:H{LAST})",number_format='0.000" kPa"')
    merge(ws,"A35:J37","Kontroller at rette lengder fordeles til riktige diametre. Tabellen summerer friksjon; reduksjonenes diameterekjede og netto endehastighet styres av Reduksjoner-fanen. En aktiv diameter som ikke forekommer i denne kjeden blir avvist, slik at feil endehastighet ikke brukes. Vertikale rette meter skal også med her; høydefeltene beregner bare hydrostatisk bidrag.",GREY,PALE,10)
    ws.print_area="A1:J37"


def valve_catalog(wb):
    if __package__:
        from .generic_valves import model_data
    else:
        from generic_valves import model_data
    general = model_data()
    path=ROOT/"data"/"fagerberg_valves.csv"
    products=[]
    if path.exists():
        with path.open(encoding="utf-8",newline="") as file:
            products=list(csv.DictReader(file))
    ws=base_sheet(wb,"Ventilkatalog","VENTILKATALOG – GENERELLE MODELLER OG EGNE DATA","Generelle K-estimater finnes i Ventildata. Egen Δp, Kv eller K kan overstyre.",[58,30,38,85,30])
    state=f'{len(general["models"])} generelle ventilmodeller tilgjengelige. Modellgrunnlag: Crane via fluids; ikke verifisert snitt fra flere leverandører. Fagerberg-produktdata: {len(products)} importerte produkter.'
    merge(ws,"A4:E5",state,INK,AMBER,11,True)
    headers=["Valg i ventilfanen","Kategori","Produsent / modell","Kilde-URL","Hentedato"]
    for col,label in enumerate(headers,1):styled_cell(ws,f"{chr(64+col)}6",label)
    generic=[("Egen ventil (dokumentert)","Egen","","",""),("Kuleventil","Generisk","","",""),("Reguleringsventil","Generisk","","",""),("Tilbakeslagsventil","Generisk","","",""),("Spjeldventil","Generisk","","",""),("Sluseventil","Generisk","","","")]
    records=generic+[(m["name"],m["category"],"Modellestimat, ikke leverandørprodukt",general["sources"][0],general["retrieved"]) for m in general["models"]]+[(p["name"],p.get("category",""),p.get("manufacturer_model",""),p["url"],p.get("retrieved","")) for p in products]
    for r,record in enumerate(records,7):
        for col,value in enumerate(record,1):styled_cell(ws,f"{chr(64+col)}{r}",value)
        ws.row_dimensions[r].height=34
    end=6+len(records)
    named(wb,"Ventilvalg",f"'Ventilkatalog'!$A$7:$A${end}")
    named(wb,"Ventiloppslag",f"'Ventilkatalog'!$A$7:$E${end}")
    ws.auto_filter.ref=f"A6:E{end}"
    ws.print_area=f"A1:E{end}"
    ws.page_setup.fitToHeight=0
    return products


def fitting_diameters(wb):
    """Allow bends and own K values to use the velocity at their actual bore."""
    ws=wb["Bend"]
    ws.column_dimensions["L"].width=21
    styled_cell(ws,"L6","Egen lokal ID [mm]")
    for r in range(FIRST,LAST+1):
        styled_cell(ws,f"L{r}",None,"input","0.000","Tomt = hoved-ID. Ved bend etter reduksjon oppgis lokal ID. Eget K skal referere til denne hastigheten.")
        diameter=f'IF(L{r}="",\'Trykkberegning\'!$E$12,L{r})'
        styled_cell(ws,f"M{r}",f'=IFERROR(Beregning!$B$14*\'Trykkberegning\'!$E$12/({diameter}),"")')
        rough=f'(\'Trykkberegning\'!$E$15/({diameter}))'
        styled_cell(ws,f"N{r}",f'=IFERROR(IF(M{r}=0,0,IF(M{r}<2300,64/M{r},(-1.8*LOG10(({rough}/3.7)^1.11+6.9/M{r}))^(-2))),"")')
        for col in range(15,21):
            letter,prev=chr(64+col),chr(63+col)
            styled_cell(ws,f"{letter}{r}",f'=IFERROR(IF(M{r}=0,0,IF(M{r}<2300,N{r},(-2*LOG10({rough}/3.7+2.51/(M{r}*SQRT({prev}{r}))))^(-2))),"")')
        for col in range(13,21):ws.column_dimensions[chr(64+col)].hidden=True
        ws[f"E{r}"]=ws[f"E{r}"].value.replace("Beregning!$B$23",f"T{r}").replace('Beregning!$B$3="OK"',f'AND(Beregning!$B$3="OK",OR(M{r}=0,M{r}>=4000))')
        original=ws[f"J{r}"].value[1:].replace("Beregning!$B$14",f"M{r}")
        checks=[
            (f'OR(NOT(ISNUMBER({diameter})),{diameter}<=0,\'Trykkberegning\'!$E$15>={diameter})',"FEIL: lokal diameter / ruhet"),
            (f'AND(ABS(({diameter})-\'Trykkberegning\'!$E$12)>0.000001,COUNTIFS(Reduksjoner!$B${FIRST}:$B${LAST},1,Reduksjoner!$D${FIRST}:$D${LAST},{diameter})=0)',"FEIL: lokal ID mangler i reduksjonskjeden"),
        ]
        ws[f"J{r}"]=f'=IFERROR(IF(OR(C{r}="",C{r}=0),"Ikke med",{condition_chain(checks,original)}),"FEIL: lokal diameter")'
        ws[f"H{r}"]=guarded(f'G{r}*Beregning!$B$25*(\'Trykkberegning\'!$E$12/({diameter}))^4/1000',f'J{r}="OK"')
    dv_number(ws,f"L{FIRST}:L{LAST}")
    ws.print_area="A1:L36"


def valve_options(wb):
    ws=wb["Ventiler"]
    ws.unmerge_cells("A1:F2");merge(ws,"A1:M2","VENTILER – MODELLVALG OG TRYKKFALL",WHITE,NAVY,22,True)
    ws.unmerge_cells("A4:F4");merge(ws,"A4:M4","Velg modell, antall og dokumentert Δp per enhet. Alternativt oppgis faktisk Kv, DN og åpning; kv-beregning må bekreftes for valgt væske.",INK,BLUE,10)
    ws.unmerge_cells("A5:F5");merge(ws,"A5:M5","Eget Δp i kPa overstyrer Kv. Kv: Δp [kPa] = 100 · (ρ/1000) · (Q [m³/h]/Kv)². Ingen automatisk viskositets- eller kavitasjonskorreksjon.",INK,AMBER,10)
    for col,width in {"G":14,"H":18,"I":18,"J":20,"K":65,"L":25,"M":23}.items():ws.column_dimensions[col].width=width
    for col,label in {"G":"DN / størrelse","H":"Dokumentert Kv [m³/h]","I":"Åpning [%]","J":"Δp brukt [kPa]","K":"Produktlenke","L":"Kv-modell","M":"Kv verifisert?"}.items():styled_cell(ws,f"{col}6",label)
    dv_list(ws,"Ventilvalg",f"A{FIRST}:A{LAST}")
    dv_list(ws,"NominelleDN",f"G{FIRST}:G{LAST}")
    dv_list(ws,'"Nei,Ja"',f"M{FIRST}:M{LAST}")
    dv_number(ws,f"H{FIRST}:H{LAST}")
    dv_number(ws,f"I{FIRST}:I{LAST}",0,100)
    for r in range(FIRST,LAST+1):
        for col,value in {"G":None,"H":None,"I":100,"M":"Nei"}.items():styled_cell(ws,f"{col}{r}",value,"input","0.000")
        styled_cell(ws,f"J{r}",f'=IFERROR(IF(C{r}<>"",C{r},IF(AND(H{r}>0,ISNUMBER(H{r})),100*(Beregning!$B$10/1000)*(\'Trykkberegning\'!$E$25/H{r})^2,"")),"")',number_format="0.000")
        styled_cell(ws,f"K{r}",f'=IFERROR(VLOOKUP(A{r},Ventiloppslag,4,FALSE),"")')
        styled_cell(ws,f"L{r}",f'=IF(C{r}<>"","Eget trykkfall","Kv ved angitt åpning")')
        checks=[
            (f'OR(B{r}="",B{r}=0)',"Ikke med"),
            (f'OR(NOT(ISNUMBER(B{r})),B{r}<0,B{r}<>INT(B{r}))',"FEIL: antall"),
            (f'A{r}=""',"FEIL: angi ventil"),
            (f'AND(C{r}<>"",OR(NOT(ISNUMBER(C{r})),C{r}<0))',"FEIL: trykkfall"),
            (f'AND(C{r}="",OR(NOT(ISNUMBER(H{r})),H{r}<=0,G{r}="",NOT(ISNUMBER(I{r})),I{r}<=0,I{r}>100,E{r}="",M{r}<>"Ja"))',"FEIL: dokumenter Kv, DN, åpning og medium"),
        ]
        styled_cell(ws,f"F{r}",f'=IFERROR({condition_chain(checks)},"FEIL: kontroller tall")')
        styled_cell(ws,f"D{r}",guarded(f'B{r}*J{r}',f'F{r}="OK"'),number_format="0.000")
    ws.print_area="A1:M36"


def integrate(wb):
    main,calc=wb["Trykkberegning"],wb["Beregning"]
    merge(main,"A34:D34","Rette rør: beregningsmetode",size=10)
    styled_cell(main,"E34","Hoveddimensjon","input")
    main.merge_cells("E34:F34")
    main.row_dimensions[34].height=30
    dv_list(main,'"Hoveddimensjon,Rørstrekk-tabell"',"E34")
    calc["B29"]=guarded(f'IF(\'Trykkberegning\'!E34="Rørstrekk-tabell",SUM(Rørstrekk!H{FIRST}:H{LAST})*1000,B23*\'Trykkberegning\'!E29/B8*B25)','B4="OK"')
    calc["B34"]=guarded(f'SUM(Reduksjoner!J{FIRST}:J{LAST})*1000','B4="OK"')
    styled_cell(calc,"A34","Reduksjoner: irreversibelt tap")
    styled_cell(calc,"C34","Pa")
    calc["B35"]=guarded('SUM(B29:B34)','B4="OK"')
    styled_cell(calc,"A40","Hastighetsbidrag ved endene")
    styled_cell(calc,"B40",guarded('B10/2*((B7/(PI()*(Reduksjoner!F4/1000)^2/4))^2-B13^2)','B4="OK"'),number_format="0.000000")
    styled_cell(calc,"C40","Pa")
    calc["B37"]=guarded('B35+B36+B40','B4="OK"')
    # Wrap the original complete checks; new tables must also validate.
    old=calc["B4"].value[1:]
    checks=[
        ('AND(\'Trykkberegning\'!E34<>"Hoveddimensjon",\'Trykkberegning\'!E34<>"Rørstrekk-tabell")',"FEIL: velg beregningsmetode for rette rør."),
        (f'COUNTIF(Reduksjoner!K{FIRST}:K{LAST},"OK")+COUNTIF(Reduksjoner!K{FIRST}:K{LAST},"Ikke med")<>{LAST-FIRST+1}',"FEIL: kontroller Reduksjoner."),
        (f'AND(\'Trykkberegning\'!E34="Rørstrekk-tabell",COUNTIF(Rørstrekk!I{FIRST}:I{LAST},"OK")+COUNTIF(Rørstrekk!I{FIRST}:I{LAST},"Ikke med")<>{LAST-FIRST+1})',"FEIL: kontroller Rørstrekk."),
    ]
    extra=condition_chain(checks)
    calc["B4"]=f'=IFERROR(IF(({old})<>"OK",({old}),{extra}),"FEIL: kontroller nye inndata.")'
    # 29 original chart header; add reducer and velocity contributions at the end.
    for r,label,ref in [(36,"Reduksjoner",34),(37,"Hastighetsbidrag",40)]:
        styled_cell(calc,f"D{r}",label)
        styled_cell(calc,f"E{r}",f'=IF(B4="OK",B{ref}/100000,0)',number_format="0.000")
    for r in range(30,38):
        label=calc[f"D{r}"].value
        calc[f"D{r}"]=f'=IF(AND(B4="OK",E{r}<>0),"{label}","")'
        calc[f"E{r}"].number_format='0.000" bar";-0.000" bar";""'
    chart=main._charts[0]
    chart.series[0].val.numRef.f="'Beregning'!$E$30:$E$37"
    chart.series[0].cat.numRef.f="'Beregning'!$D$30:$D$37"
    chart.x_axis.tickLblPos="none"
    chart.x_axis.delete=True
    # In a horizontal bar chart, y_axis is the horizontal numeric axis.
    # Keep the unit in its title and avoid inheriting long bar-label formats.
    chart.y_axis.numFmt=NumFmt(formatCode="0.###",sourceLinked=False)
    chart.y_axis.txPr=RichText(bodyPr=RichTextProperties(rot=-45*60000),p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=900,solidFill=INK)))])
    chart.dLbls=DataLabelList(dLblPos="outEnd",showVal=True,showCatName=True,showSerName=False,showLegendKey=False,showPercent=False,separator="\n",numFmt='0.000" bar";-0.000" bar";""')
    chart.layout=Layout(manualLayout=ManualLayout(x=.23,y=.18,w=.54,h=.68,xMode="factor",yMode="factor",wMode="factor",hMode="factor",layoutTarget="inner"))
    chart.height=12.0
    main["H10"]="Tap + høydebidrag + netto hastighetsbidrag ved endene"
    main["H14"]="Referansehastighet (hoved-ID)"
    main["H15"]="Reynolds (hoved-ID)"
    main["H16"]="Regime (hoved-ID)"
    main["H17"]="Darcy f (hoved-ID)"
    # Links to additions, keeping the existing main inputs and cached-output addresses.
    for r,name,label in [(42,"Reduksjoner","Reduksjoner og endehastighet →"),(43,"Rørstrekk","Rette rør med flere diametre →"),(44,"Rørstyrke","Styrkeberegning og sikkerhetsfaktor →"),(45,"Ventildata","Generelle ventiler: K, Kv og datakilder →")]:
        c=merge(main,f"A{r}:F{r}",label,TEAL,BLUE,11,True)
        c.hyperlink=f"#'{name}'!A1"
    main.unmerge_cells("A40:N41")
    merge(main,"A40:F41","MODELL: seriell strømningsvei for Newtonsk væske. Ulike diametre: Reduksjoner/Rørstrekk. Rørstyrke: separat kontroll av indre overtrykk. Se Veiledning.",GREY,PALE,10)
    main.print_area="A1:N45"
    calc.print_area="A1:E40"


def add_extensions(wb):
    strength(wb)
    reducers(wb)
    pipe_sections(wb)
    fitting_diameters(wb)
    valve_catalog(wb)
    valve_options(wb)
    if __package__:
        from .generic_valves import add_generic_valves
    else:
        from generic_valves import add_generic_valves
    add_generic_valves(wb)
    integrate(wb)
