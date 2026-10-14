"""Auditable full-open valve estimates; no claimed manufacturer averages."""
import json
import math

from openpyxl.utils import get_column_letter

if __package__:
    from .build_pressure_workbook import ROOT, DN_OD, FIRST, LAST, WHITE, TEAL, INK, BLUE, AMBER, base_sheet, merge, styled_cell, dv_number, named
    from .pressure_extensions import condition_chain, guarded
else:
    from build_pressure_workbook import ROOT, DN_OD, FIRST, LAST, WHITE, TEAL, INK, BLUE, AMBER, base_sheet, merge, styled_cell, dv_number, named
    from pressure_extensions import condition_chain, guarded


def model_data():
    return json.loads((ROOT / "data/generic_valve_models.json").read_text(encoding="utf-8"))


def crane_reference_friction(diameter):
    """Solve the documented fluids fit to Crane's tabulated reference fT.

    This is a fitting reference factor, independent of the actual fluid/pipe f.
    """
    reynolds = 7.5e6 * diameter
    relative_roughness = 3.4126825352925e-5 * diameter ** -1.0112
    lo, hi = .001, .2
    for _ in range(90):
        f = (lo + hi) / 2
        residual = 1 / math.sqrt(f) + 2 * math.log10(relative_roughness / 3.7 + 2.51 / (reynolds * math.sqrt(f)))
        if residual > 0:
            lo = f
        else:
            hi = f
    return (lo + hi) / 2


def add_generic_valves(wb):
    data = model_data()
    models = data["models"]
    table = base_sheet(wb, "Ventildata", "GENERELLE VENTILER – K OG EKVIVALENT Kv",
                       "Full åpning · modellestimater · DN15–300 · faktisk lokal ID brukes i ventilberegningen",
                       [58,48,14,19,15,18,17,24,90,85])
    merge(table, "A4:J5", "Ingen verifisert snittverdi fra flere leverandører. Crane-referanser gjengitt i fluids; skyvespjeld bruker et uttrykkelig sluseanslag. Referanse-ID i denne tabellen = DN i mm, ikke produsentens ventilboring. Små spjeld utenfor modellområdet har ingen verdi.", INK, AMBER, 11, True)
    table.row_dimensions[4].height = 40
    table.row_dimensions[5].height = 32
    headers = ["Oppslagsnøkkel", "Ventilvalg", "DN", "Referanse-ID [mm]", "N i K=N·fT", "fT referanse", "K referanse", "Kv ekv. [m³/h]", "Grunnlag / gyldighet", "Kilde"]
    for col, label in enumerate(headers, 1):
        styled_cell(table, f"{chr(64+col)}6", label)
    row = 7
    for model in models:
        for dn, _ in DN_OD:
            valid = dn >= model["minimum_dn"]
            factor = model["multiplier_small"] if dn <= 200 else model["multiplier_large"]
            f = crane_reference_friction(dn / 1000)
            k = factor * f
            kv = 3600 * math.pi * (dn / 1000) ** 2 / 4 * math.sqrt(200 / k)
            record = [f'{model["name"]}|DN{dn}', model["name"], f"DN{dn}", dn,
                      factor if valid else None, f if valid else None, k if valid else None,
                      kv if valid else None, model["basis"] if valid else "Utenfor modellområde: bruk dokumentert Kv eller eget K.", data["sources"][0]]
            for col, value in enumerate(record, 1):
                styled_cell(table, f"{chr(64+col)}{row}", value, number_format="0.000000" if col in (6,7) else "0.000")
            table.row_dimensions[row].height = 55
            row += 1
    named(wb, "GenerelleVentildata", f"'Ventildata'!$A$7:$J${row-1}")
    table.auto_filter.ref = f"A6:J{row-1}"
    table.print_area = f"A1:J{row+5}"
    table.page_setup.fitToHeight = 0
    merge(table, f"A{row+1}:J{row+3}", "Δp = K·ρ·v²/2. K er dimensjonsløs og refererer til faktisk lokal rør-ID. Ekvivalent Kv = 3600·A·√(200/K), med A i m². fT: Colebrook med Re*=7,5·10⁶·D og ε/D=3,4126825352925·10⁻⁵·D^(−1,0112), D i meter (fluids-tilpasning til Crane). Dette er ikke rørets friksjonsfaktor. Generell modell brukes bare ved Re≥4000 eller null strømning, og 100 % åpning.", INK, BLUE, 10)
    merge(table, f"A{row+4}:J{row+5}", f'Kilde innhentet {data["retrieved"]}: {data["sources"][1]}', size=10)

    ws = wb["Ventiler"]
    ws.unmerge_cells("A4:M4")
    merge(ws, "A4:R4", "Generell ventil + antall gir automatisk trykktap. DN følger hovedfanen; bruk lokal ID etter reduksjon. Estimatene gjelder helt åpen ventil og turbulent strømning.", INK, BLUE, 10)
    ws.unmerge_cells("A5:M5")
    merge(ws, "A5:R5", "Prioritet: eget Δp → dokumentert Kv → eget K → generell K-modell. K og ekvivalent Kv vises. Skyvespjeld er et sluseanslag; ingen verifisert leverandørgjennomsnitt eller åpningstilpasning.", INK, AMBER, 10)
    for col,width in {"N":16,"O":17,"P":19,"Q":19,"R":24}.items():
        ws.column_dimensions[col].width = width
    for col,label in {"H":"Eget dokumentert Kv [m³/h]", "L":"Metode / estimatgrunnlag", "M":"Eget Kv verifisert?", "N":"Eget K", "O":"K brukt, lokal ID", "P":"Egen lokal ID [mm]", "Q":"ID brukt [mm]", "R":"Kv ekv. fra K [m³/h]"}.items():
        styled_cell(ws, f"{col}6", label)
    ws.column_dimensions["L"].width = 46
    ws.column_dimensions["F"].width = 46
    ws.row_dimensions[6].height = 55
    ws.freeze_panes = "C7"
    # Keep existing manual cases and generic names available for old workbooks.
    for r in range(FIRST,LAST+1):
        ws.row_dimensions[r].height = 42
        if r - FIRST < len(models):
            ws[f"A{r}"] = models[r-FIRST]["name"]
        styled_cell(ws,f"G{r}","='Trykkberegning'!E7", "input")
        styled_cell(ws,f"N{r}",None,"input","0.000000", "Eget K referert til hastighet ved ID brukt. Dokumenter kilde i kommentar. Kan også brukes ved andre åpninger/strømningsregimer når K gjelder disse.")
        styled_cell(ws,f"P{r}",None,"input","0.000", "Tomt = hoved-ID. Etter reduksjon oppgis aktuell ID som finnes i Reduksjoner. DN i G velges separat; DN er nominelt, ikke ventilens borediameter.")
        styled_cell(ws,f"Q{r}",f'=IF(P{r}="",\'Trykkberegning\'!$E$12,P{r})',number_format="0.000")
        # Solve reference Colebrook in Excel so wall/ID changes update K too.
        d=f"(Q{r}/1000)"
        re=f"(7.5E6*{d})"
        ed=f"(3.4126825352925E-5*{d}^(-1.0112))"
        styled_cell(ws,f"S{r}",f'=IFERROR((-1.8*LOG10(({ed}/3.7)^1.11+6.9/{re}))^(-2),"")')
        for col in range(20,26):
            letter,prev=chr(64+col),chr(63+col)
            styled_cell(ws,f"{letter}{r}",f'=IFERROR((-2*LOG10({ed}/3.7+2.51/({re}*SQRT({prev}{r}))))^(-2),"")')
        styled_cell(ws,f"Z{r}",f'=IFERROR(VLOOKUP(A{r}&"|"&G{r},GenerelleVentildata,5,FALSE),0)')
        styled_cell(ws,f"AA{r}",f'=IFERROR(Beregning!$B$14*\'Trykkberegning\'!$E$12/Q{r},"")')
        # Main-series DN is authoritative at the inlet. A reduced section's
        # DN cannot be inferred from its ID: wall thickness is not known there.
        styled_cell(ws,f"AB{r}",f'=IFERROR(IF(AND(ABS(Q{r}-\'Trykkberegning\'!$E$12)<0.000001,\'Trykkberegning\'!$E$6<>"Metrisk OD"),\'Trykkberegning\'!$E$7,""),"")')
        styled_cell(ws,f"O{r}",f'=IFERROR(IF(C{r}<>"","",IF(H{r}<>"",200*(3600*PI()*(Q{r}/1000)^2/4/H{r})^2,IF(N{r}<>"",N{r},IF(Z{r}>0,Z{r}*Y{r},"")))),"")',number_format="0.000000")
        styled_cell(ws,f"R{r}",f'=IFERROR(IF(O{r}>0,3600*PI()*(Q{r}/1000)^2/4*SQRT(200/O{r}),""),"")',number_format="0.000")
        styled_cell(ws,f"J{r}",f'=IFERROR(IF(C{r}<>"",C{r},IF(H{r}<>"",100*(Beregning!$B$10/1000)*(\'Trykkberegning\'!$E$25/H{r})^2,IF(ISNUMBER(O{r}),O{r}*Beregning!$B$25*(\'Trykkberegning\'!$E$12/Q{r})^4/1000,""))),"")',number_format="0.000")
        styled_cell(ws,f"L{r}",f'=IF(C{r}<>"","Eget trykkfall",IF(H{r}<>"","Dokumentert Kv",IF(N{r}<>"","Eget K",IF(Z{r}<=0,"Ingen generell modell",IF(A{r}="{models[1]["name"]}","ANSLAG: skyvespjeld fra sluse","Modellestimat: Crane, full åpning")))))')
        checks=[
            (f'OR(B{r}="",B{r}=0)',"Ikke med"),
            (f'OR(NOT(ISNUMBER(B{r})),B{r}<0,B{r}<>INT(B{r}))',"FEIL: antall"),
            (f'A{r}=""',"FEIL: angi ventil"),
            (f'AND(C{r}<>"",OR(NOT(ISNUMBER(C{r})),C{r}<0))',"FEIL: trykkfall"),
            (f'AND(C{r}="",H{r}<>"",OR(NOT(ISNUMBER(H{r})),H{r}<=0,G{r}="",NOT(ISNUMBER(I{r})),I{r}<=0,I{r}>100,E{r}="",M{r}<>"Ja"))',"FEIL: dokumenter Kv, DN, åpning og medium"),
            (f'AND(C{r}="",H{r}="",OR(NOT(ISNUMBER(Q{r})),Q{r}<=0))',"FEIL: lokal ID"),
            (f'AND(C{r}="",H{r}="",ABS(Q{r}-\'Trykkberegning\'!$E$12)>0.000001,COUNTIFS(Reduksjoner!$B${FIRST}:$B${LAST},1,Reduksjoner!$D${FIRST}:$D${LAST},Q{r})=0)',"FEIL: lokal ID mangler i reduksjonskjeden"),
            (f'AND(C{r}="",H{r}="",N{r}<>"",OR(NOT(ISNUMBER(N{r})),N{r}<0,E{r}="",NOT(ISNUMBER(I{r})),I{r}<=0,I{r}>100))',"FEIL: dokumenter eget K og åpning"),
            (f'AND(C{r}="",H{r}="",N{r}="",OR(Z{r}<=0,AND(AB{r}<>"",G{r}<>AB{r})))',"FEIL: ingen generell modell for DN / lokal ID"),
            (f'AND(C{r}="",H{r}="",N{r}="",I{r}<>100)',"FEIL: generell modell krever 100 % åpning"),
            (f'AND(C{r}="",H{r}="",N{r}="",AA{r}>0,AA{r}<4000)',"FEIL: generell modell krever Re >=4000"),
        ]
        styled_cell(ws,f"F{r}",f'=IFERROR({condition_chain(checks)},"FEIL: kontroller ventilinndata")')
        styled_cell(ws,f"D{r}",guarded(f'B{r}*J{r}',f'F{r}="OK"'),number_format="0.000")
    for col in range(19,29):
        ws.column_dimensions[get_column_letter(col)].hidden=True
    dv_number(ws,f"N{FIRST}:N{LAST}",0)
    dv_number(ws,f"P{FIRST}:P{LAST}",1e-9)
    ws.auto_filter.ref=f"A6:R{LAST}"
    ws.print_area="A1:R36"
