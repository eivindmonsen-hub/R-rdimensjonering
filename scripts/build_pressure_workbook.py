"""Build a Norwegian, macro-free liquid pipe pressure-loss workbook."""
from __future__ import annotations

import argparse
from copy import copy
import os
from pathlib import PurePosixPath
import re
import signal
import shutil
import subprocess
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
FIRST, LAST = 7, 31
NAVY, TEAL, BLUE = "142D4E", "087F8C", "E5F2FF"
PALE, WHITE, INK, GREY = "F1F5F9", "FFFFFF", "20354C", "617589"
AMBER, RED, GREEN = "FFF0CE", "FFE5E5", "E1F3EB"
EDGE = Side(style="thin", color="D6E1EA")

SERIES = ["EN 10220 / ISO 4200", "ISO 1127", "Metrisk OD"]
DN_OD = [(15, 21.3), (20, 26.9), (25, 33.7), (32, 42.4), (40, 48.3),
         (50, 60.3), (65, 76.1), (80, 88.9), (100, 114.3), (125, 139.7),
         (150, 168.3), (200, 219.1), (250, 273.0), (300, 323.9)]
METRIC_OD = [20, 25, 32, 40, 50, 63, 75, 90, 110, 125, 140, 160, 180, 200,
             225, 250, 280, 315]
MATERIALS = [
    ("P235GH – karbonstål", "P235GH", 0.045, "|EN 10216-2|EN 10217-2|EN 10217-5|", "Sømløst/sveiset trykkrør for forhøyet temperatur."),
    ("P265GH – karbonstål", "P265GH", 0.045, "|EN 10216-2|EN 10217-2|EN 10217-5|", "Sømløst/sveiset trykkrør for forhøyet temperatur."),
    ("P235TR2 – karbonstål", "P235TR2", 0.045, "|EN 10216-1|EN 10217-1|", "Trykkrør med spesifiserte egenskaper ved romtemperatur."),
    ("1.4307 / 304L – rustfritt", "X2CrNi18-9 / 1.4307", 0.015, "|EN 10216-5|EN 10217-7|", "Austenittisk rustfritt trykkrør; EN 10088-1 materialbetegnelse."),
    ("1.4404 / 316L – rustfritt", "X2CrNiMo17-12-2 / 1.4404", 0.015, "|EN 10216-5|EN 10217-7|", "Austenittisk rustfritt trykkrør; EN 10088-1 materialbetegnelse."),
    ("1.4462 / 2205 – duplex", "X2CrNiMoN22-5-3 / 1.4462", 0.015, "|EN 10216-5|EN 10217-7|", "Duplex trykkrør; EN 10088-1 materialbetegnelse."),
]
EN_STANDARDS = ["EN 10216-1", "EN 10217-1", "EN 10216-2", "EN 10217-2",
                "EN 10217-5", "EN 10216-5", "EN 10217-7", "Annen / dokumentert"]
MEDIA = [
    ("Vann 5 °C", 999.97, 1.519, "Omtrentlige verdier ved atmosfæretrykk."),
    ("Vann 20 °C", 998.2, 1.002, "Omtrentlige verdier ved atmosfæretrykk."),
    ("Vann 40 °C", 992.2, 0.653, "Omtrentlige verdier ved atmosfæretrykk."),
    ("Vann 60 °C", 983.2, 0.466, "Omtrentlige verdier ved atmosfæretrykk."),
    ("Vann 80 °C", 971.8, 0.354, "Omtrentlige verdier ved atmosfæretrykk."),
    ("Olje ISO VG 32, 40 °C", 850.0, 27.2, "Nominell ν=32 mm²/s; antatt ρ=850 kg/m³. Bruk leverandørdata."),
    ("Eget medium", 0.0, 0.0, "Begge egenskaper må fylles inn i feltene for egne verdier."),
]
FITTINGS = [
    ("Bend, lang radius", 1.5, "Rennels: rundet bend. Startanslag R/Di=1,5; oppgi faktisk radiusforhold."),
    ("Bend, kort radius", 1.0, "Rennels: rundet bend. Startanslag R/Di=1,0; oppgi faktisk radiusforhold."),
    ("Y-stykke, gjennomløp", 0, "Oppgi dokumentert K for valgt strømningsvei og fordeling."),
    ("Y-stykke, avgrening", 0, "Oppgi dokumentert K for valgt strømningsvei og fordeling."),
    ("T-stykke, gjennomløp", 0, "Oppgi dokumentert K for valgt strømningsvei og fordeling."),
    ("T-stykke, avgrening", 0, "Oppgi dokumentert K for valgt strømningsvei og fordeling."),
    ("Egen komponent", 0, "Oppgi eget K med referanse til hastigheten i denne rørstrekningen."),
]


def merge(ws, area, value, color=INK, fill=None, size=11, bold=False):
    ws.merge_cells(area)
    for row in ws[area]:
        for cell in row:
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    cell = ws[area.split(":")[0]]
    cell.value = value
    cell.font = Font(name="Calibri", size=size, bold=bold, color=color)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    return cell


def base_sheet(wb, name, title, subtitle, widths):
    ws = wb.create_sheet(name)
    end = chr(64 + len(widths))
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = TEAL
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=.25, right=.25, top=.35, bottom=.35, header=.15, footer=.15)
    ws.oddFooter.center.text = "R-rdimensjonering | &A | side &P av &N"
    ws.oddFooter.center.size = 8
    ws.freeze_panes = "A7"
    for row in range(1, 55):
        ws.row_dimensions[row].height = 23
    merge(ws, f"A1:{end}2", title, WHITE, NAVY, 22, True)
    merge(ws, f"A3:{end}3", subtitle, GREY, size=10)
    return ws


def styled_cell(ws, address, value=None, kind="calc", number_format=None, note=None):
    cell = ws[address]
    cell.value = value
    cell.fill = PatternFill("solid", fgColor=BLUE if kind == "input" else PALE)
    cell.font = Font(name="Calibri", size=11, color="086AA5" if kind == "input" else INK)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = Border(bottom=EDGE)
    cell.protection = Protection(locked=(kind != "input"))
    if number_format:
        cell.number_format = number_format
    if note:
        cell.comment = Comment(note, "R-rdimensjonering")
    return cell


def dv_list(ws, formula, cells, title="Velg verdi"):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle, dv.error = "Ugyldig valg", "Velg fra listen, eller bruk et felt for egne verdier."
    dv.showInputMessage, dv.promptTitle = True, title
    dv.prompt = "Bruk nedtrekkslisten. Blå felt kan redigeres."
    dv.showDropDown = False
    ws.add_data_validation(dv)
    dv.add(cells)


def dv_number(ws, cells, minimum=0, maximum=1e9, integer=False, allow_blank=True):
    dv = DataValidation(type="whole" if integer else "decimal", operator="between",
                        formula1=minimum, formula2=maximum, allow_blank=allow_blank)
    dv.showErrorMessage = True
    dv.errorTitle, dv.error = "Ugyldig tall", f"Oppgi et {'heltall' if integer else 'tall'} mellom {minimum} og {maximum}."
    ws.add_data_validation(dv)
    dv.add(cells)


def named(wb, name, target):
    wb.defined_names.add(DefinedName(name, attr_text=target))


def references(wb):
    ws = base_sheet(wb, "Rørdata", "RØRSERIER OG DIMENSJONER", "Utvalg vanlige ytre diametre. DN er nominelt; innvendig diameter krever veggtykkelse eller egen verdi.", [45, 28, 22, 18, 65])
    ws.append([])
    for col, label in enumerate(["Oppslagsnøkkel", "Rørserie", "DN / OD", "Ytre diameter [mm]", "Merknad / normreferanse"], 1):
        styled_cell(ws, f"{chr(64+col)}6", label)
    records = []
    for series in SERIES[:2]:
        for dn, od in DN_OD:
            records.append((f"{series}|DN{dn}", series, f"DN{dn}", od,
                            "Vanlig OD-utvalg, ikke komplett normtabell. Kontroller mot rørsertifikat og aktuell normutgave."))
    for od in METRIC_OD:
        records.append((f"Metrisk OD|OD{od}", "Metrisk OD", f"OD{od}", od,
                        "Metrisk OD-valg i mm; ingen produktstandard eller veggserie er implisert."))
    for r, record in enumerate(records, 7):
        for col, value in enumerate(record, 1):
            styled_cell(ws, f"{chr(64+col)}{r}", value, number_format="0.0" if col == 4 else None)
        ws.row_dimensions[r].height = 30
    named(wb, "Roroppslag", f"'Rørdata'!$A$7:$E${6+len(records)}")
    ws.auto_filter.ref = f"A6:E{6+len(records)}"
    ws.print_area = f"B1:E{6+len(records)}"
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:6"

    ws = base_sheet(wb, "Materialer", "MATERIALER OG EN-PRODUKTSTANDARDER", "Produktstandard velges separat. Ingen beregning av tillatt arbeidstrykk, spenninger eller godstykkelse.", [37, 33, 20, 45, 65])
    for col, label in enumerate(["Materialvalg", "Materialbetegnelse", "Ruhet, startanslag [mm]", "Tillatte valg i denne malen", "Beskrivelse"], 1):
        styled_cell(ws, f"{chr(64+col)}6", label)
    for r, record in enumerate(MATERIALS, 7):
        for col, value in enumerate(record, 1):
            styled_cell(ws, f"{chr(64+col)}{r}", value, number_format="0.000" if col == 3 else None)
        ws.row_dimensions[r].height = 45
    named(wb, "Materialdata", "'Materialer'!$A$7:$E$12")
    merge(ws, "A15:E16", "Ruhet er et generelt startanslag, ikke en egenskap garantert av EN-standarden. Overstyr med dokumentert ruhet for faktisk overflate, sveiser, belegg og aldring.", INK, AMBER)
    merge(ws, "A18:E19", "EN 10216: sømløse trykkrør. EN 10217: sveisede trykkrør. Relevant del og materialkvalitet må bekreftes mot leverandørens dokumentasjon. «Annen / dokumentert» tillater et særskilt dokumentert valg.", INK, PALE)
    ws.print_area = "A1:E20"

    ws = base_sheet(wb, "Medier", "MEDIER OG REFERANSEVERDIER", "Tetthet og dynamisk viskositet ved angitt temperatur; overstyr ved faktisk driftsbetingelse.", [34, 22, 26, 90])
    for col, label in enumerate(["Medium", "Tetthet [kg/m³]", "Dynamisk viskositet [mPa·s]", "Betingelser / merknad"], 1):
        styled_cell(ws, f"{chr(64+col)}6", label)
    for r, record in enumerate(MEDIA, 7):
        for col, value in enumerate(record, 1):
            styled_cell(ws, f"{chr(64+col)}{r}", value, number_format="0.000" if col in (2, 3) else None)
        ws.row_dimensions[r].height = 34
    named(wb, "Mediedata", "'Medier'!$A$7:$D$13")
    merge(ws, "A16:D17", "Vannverdiene er avrundede referanseverdier i tråd med vanlige tabeller basert på IAPWS. De er ikke en implementasjon av temperatur- og trykkavhengige tilstandsligninger. VG 32 angir nominell kinematisk viskositet ved 40 °C, ikke en universell oljeegenskap.", INK, AMBER)
    ws.print_area = "A1:D18"

    ws = wb.create_sheet("_Lister")
    lists = {
        "Rorserier": SERIES,
        "NominelleDN": [f"DN{dn}" for dn, _ in DN_OD] + ["Egendefinert"],
        "MetriskeOD": [f"OD{od}" for od in METRIC_OD] + ["Egendefinert"],
        "Materialvalg": [m[0] for m in MATERIALS],
        "ENStandarder": EN_STANDARDS,
        "Medievalg": [m[0] for m in MEDIA],
        "Bendtyper": [f[0] for f in FITTINGS],
    }
    for col, (name, values) in enumerate(lists.items(), 1):
        letter = chr(64+col)
        ws.cell(1, col, name)
        for r, value in enumerate(values, 2):
            ws.cell(r, col, value)
        named(wb, name, f"'_Lister'!${letter}$2:${letter}${len(values)+1}")
    for r, item in enumerate(FITTINGS, 2):
        for col, value in enumerate(item, 9):
            ws.cell(r, col, value)
    named(wb, "Benddata", "'_Lister'!$I$2:$K$8")
    ws.sheet_state = "hidden"


def inputs(wb):
    ws = base_sheet(wb, "Trykkberegning", "TRYKKBEREGNING AV RØR", "Væsker · én rørstrekning · eksempeldata ved åpning · blå felt er inndata", [18, 11, 11, 13, 23, 13, 3, 15, 12, 12, 12, 13, 11, 12])
    ws.freeze_panes = "E6"
    for row, title in [(5, "1  RØR OG MATERIALE"), (19, "2  MEDIUM OG VOLUMSTRØM"), (28, "3  RØRLENGDE OG HØYDEFORSKJELL")]:
        merge(ws, f"A{row}:F{row}", title, WHITE, TEAL, 12, True)
    entries = {
        6: ("Rørserie / dimensjonsstandard", SERIES[0], "", "input"),
        7: ("Nominell dimensjon / metrisk OD", "DN80", "", "input"),
        8: ("Utvendig diameter, brukt", '=IF(E9<>"",E9,IFERROR(VLOOKUP(E6&"|"&E7,Roroppslag,4,FALSE),0))', "mm", "calc"),
        9: ("Egen utvendig diameter, valgfritt", None, "mm", "input"),
        10: ("Veggtykkelse", 3.2, "mm", "input"),
        11: ("Egen innvendig diameter, valgfritt", None, "mm", "input"),
        12: ("Innvendig diameter, brukt", '=IF(E11<>"",E11,IFERROR(E8-2*E10,0))', "mm", "calc"),
        13: ("Materiale", MATERIALS[0][0], "", "input"),
        14: ("EN-produktstandard", "EN 10216-2", "", "input"),
        15: ("Absolutt ruhet ε, brukt", '=IF(E16<>"",E16,IFERROR(VLOOKUP(E13,Materialdata,3,FALSE),0))', "mm", "calc"),
        16: ("Egen absolutt ruhet, valgfritt", None, "mm", "input"),
        17: ("Antall flensede rørkoblinger", 0, "stk", "input"),
        18: ("K per flenskobling", 0, "–", "input"),
        20: ("Medium / temperaturreferanse", "Vann 20 °C", "", "input"),
        21: ("Tetthet ρ, brukt", '=IF(E23<>"",E23,IFERROR(VLOOKUP(E20,Mediedata,2,FALSE),0))', "kg/m³", "calc"),
        22: ("Dynamisk viskositet μ, brukt", '=IF(E24<>"",E24,IFERROR(VLOOKUP(E20,Mediedata,3,FALSE),0))', "mPa·s", "calc"),
        23: ("Egen tetthet, valgfritt", None, "kg/m³", "input"),
        24: ("Egen dynamisk viskositet, valgfritt", None, "mPa·s", "input"),
        25: ("Volumstrøm Q", 36, "m³/h", "input"),
        26: ("Kinematisk viskositet ν, brukt", '=IFERROR(1000*E22/E21,0)', "mm²/s", "calc"),
        29: ("Rette rørmeter L, totalt", 100, "m", "input"),
        30: ("Høydemeter opp", 10, "m", "input"),
        31: ("Høydemeter ned", 2, "m", "input"),
        32: ("Netto høydeforskjell", '=IFERROR(E30-E31,0)', "m", "calc"),
        33: ("Innløpstrykk, valgfritt", None, "bar(g)", "input"),
    }
    notes = {
        6: "EN 10220 / ISO 4200 og ISO 1127: utvalg vanlig OD for metallrør. Metrisk OD er direkte valg i mm, uten automatisk veggserie.",
        7: "Bytter du til metrisk OD, velg en OD-verdi fra listen. DN alene bestemmer ikke innvendig diameter.",
        9: "Tomt felt bruker tabellens utvendige diameter. Egendefinert dimensjon trenger egen OD eller egen ID.",
        10: "Veggtykkelse må være kjent fra rørdata/sertifikat. Det foretas ingen mekanisk godstykkelsesberegning.",
        11: "Denne verdien overstyrer OD − 2 × veggtykkelse. Bruk for kjent/målt lysåpning, belegg eller avvikende veggserie.",
        14: "Velg kompatibel produktstandard for valgt materiale. Dette er dokumentasjon, ikke kontroll av arbeidstrykk iht. EN 13480.",
        16: "Tomt felt bruker materialets generelle startanslag. 0 gir hydraulisk glatt rør.",
        17: "Én kobling = ett flensepar med pakning. Kontroller flensstandard, f.eks. EN 1092-1, mot leverandørens dokumentasjon. Antall påvirker trykktap bare når K per kobling er større enn 0.",
        18: "K=0 er en idealisert kobling med lik diameter, full lysåpning og uten utstikkende pakning. Bruk dokumentert K ved innsnevring/avvik, med referanse til hastigheten i denne rørstrekningen. Alternativt registreres trykkfall i Komponenter, ikke begge steder.",
        24: "Dynamisk viskositet i mPa·s = cP. 1 Pa·s = 1000 mPa·s. Ikke legg inn kinematisk viskositet (cSt) her.",
        25: "Volumstrømmen er nødvendig for hastighet, Reynolds-tall og trykktap. Modellen bruker samme volumstrøm langs hele strekningen.",
        29: "Oppgi totale rette rørmeter, både horisontale og vertikale. Høydefeltene legger ikke til rørfriksjonslengde. Rennels-modellen inkluderer friksjon i den buede delen av bendene.",
        30: "Sum positive vertikale stigninger. Høydebidrag beregnes fra netto opp minus ned.",
        31: "Sum positive vertikale fall. Fall reduserer nødvendig trykkdifferanse; tapene forblir positive.",
        33: "Valgfritt manometertrykk. Utløpstrykk = innløpstrykk − nødvendig differanse, ved lik endehastighet. Kavitasjon/absolutt trykk vurderes ikke.",
    }
    for row, (label, value, unit, kind) in entries.items():
        merge(ws, f"A{row}:D{row}", label, size=10)
        fmt = "0" if row == 17 else "0.000" if row in (15,16,18,22,24,26) else "0.00"
        cell = styled_cell(ws, f"E{row}", value, kind, fmt, notes.get(row))
        if not unit:
            ws.merge_cells(f"E{row}:F{row}")
            ws[f"F{row}"].fill = copy(cell.fill)
            ws.row_dimensions[row].height = 30
            cell.number_format = "General"
        else:
            styled_cell(ws, f"F{row}", unit)
            ws[f"F{row}"].alignment = Alignment(vertical="center", indent=1)
    dv_list(ws, "Rorserier", "E6")
    dv_list(ws, 'INDIRECT(IF($E$6="Metrisk OD","MetriskeOD","NominelleDN"))', "E7")
    dv_list(ws, "Materialvalg", "E13")
    dv_list(ws, "ENStandarder", "E14")
    dv_list(ws, "Medievalg", "E20")
    for row in [9,10,11,16,23,24,25,29,30,31]:
        dv_number(ws, f"E{row}")
    dv_number(ws, "E17", 0, 100000, True)
    dv_number(ws, "E18")
    dv_number(ws, "E33", -1, 10000)

    merge(ws, "H5:N5", "RESULTAT  /  Δp", WHITE, NAVY, 12, True)
    merge(ws, "H6:N6", "Nødvendig trykkdifferanse  p inn − p ut", GREY, size=11)
    cell = merge(ws, "H7:N9", '=IF(Beregning!B4="OK",Beregning!B37/100000,"")', WHITE, TEAL, 30, True)
    cell.number_format = '0.000" bar"'
    merge(ws, "H10:N10", "Dissipative tap + ρ · g · (høyde opp − høyde ned)", GREY, size=10)
    results = {
        12: ("Dissipativt trykktap", '=IF(Beregning!B4="OK",Beregning!B35/100000,"")', '0.000" bar"'),
        13: ("Høydebidrag, med fortegn", '=IF(Beregning!B4="OK",Beregning!B36/100000,"")', '0.000" bar"'),
        14: ("Hastighet i røret", '=IF(Beregning!B3="OK",Beregning!B13,"")', '0.000" m/s"'),
        15: ("Reynolds-tall", '=IF(Beregning!B3="OK",Beregning!B14,"")', '#,##0'),
        16: ("Strømningsregime", '=IF(Beregning!B3="OK",Beregning!B24,"")', 'General'),
        17: ("Darcy friksjonsfaktor", '=IF(Beregning!B3="OK",Beregning!B23,"")', '0.00000'),
        18: ("Nødvendig trykkhøyde", '=IF(Beregning!B4="OK",Beregning!B37/(Beregning!B10*Beregning!B6),"")', '0.000" m væskesøyle"'),
        19: ("Beregnet utløpstrykk", '=IF(OR(Beregning!B4<>"OK",E33=""),"Ikke angitt",E33-Beregning!B37/100000)', '0.000" bar(g)"'),
    }
    for row, (label, formula, fmt) in results.items():
        merge(ws, f"H{row}:K{row}", label, size=10)
        c = merge(ws, f"L{row}:N{row}", formula, INK, PALE, 11, True)
        c.number_format = fmt
    merge(ws, "H21:N21", "KONTROLL AV INNDATA OG MODELL", WHITE, NAVY, 11, True)
    merge(ws, "H22:N24", '=IF(Beregning!B4<>"OK",Beregning!B4,IF(E14="Annen / dokumentert","OK – kontroller egen EN-standard mot leverandørdata.",IF(Beregning!B37<0,"OK – negativ differanse: høydefallet gir mer trykk enn tapene tar.","OK – alle aktive komponentrader er med.")))', INK, GREEN, 11, True)
    ws.conditional_formatting.add("H22:N24", FormulaRule(formula=['LEFT($H$22,4)="FEIL"'], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add("H22:N24", FormulaRule(formula=['LEFT($H$22,8)="ADVARSEL"'], fill=PatternFill("solid", fgColor=AMBER)))
    chart = BarChart()
    chart.type, chart.style = "bar", 10
    chart.title, chart.y_axis.title = "Bidrag til trykkdifferanse", "bar (med fortegn)"
    chart.x_axis.title = None
    chart.add_data(Reference(wb["Beregning"], min_col=5, min_row=29, max_row=35), titles_from_data=True)
    chart.set_categories(Reference(wb["Beregning"], min_col=4, min_row=30, max_row=35))
    chart.legend = None
    chart.height, chart.width = 7.1, 16.1
    ws.add_chart(chart, "H26")
    merge(ws, "A35:F35", "4  LEGG TIL KOMPONENTER I EGNE FANER", WHITE, TEAL, 11, True)
    for row, sheet, label in [(36,"Bend","Bend, Y- og T-stykker →"), (37,"Ventiler","Ventiler: antall og Δp per ventil →"), (38,"Komponenter","Andre komponenter: antall og Δp →")]:
        c = merge(ws, f"A{row}:F{row}", label, "086AA5", BLUE, 11, True)
        c.hyperlink = f"#'{sheet}'!A1"
    merge(ws, "A40:N41", "MODELL: Newtonsk, inkompressibel væske; konstant diameter og volumstrøm. Darcy–Weisbach + Colebrook (laminært: 64/Re). Hydraulikk, ikke mekanisk trykklasse. Se Veiledning for forutsetninger og kilder.", GREY, PALE, 10)
    ws.print_area = "A1:N41"
    return ws


def fitting_rows(wb):
    ws = base_sheet(wb, "Bend", "BEND, Y-STYKKER OG T-STYKKER", "25 rader kan kombineres. Antall 0 betyr at raden ikke er med.", [32, 12, 11, 14, 15, 14, 15, 18, 19, 43, 38])
    merge(ws, "A4:K4", "Bend: velg type og vinkel. R/Di kan overstyres (tomt = 1,5 for lang radius og 1,0 for kort). Eget K overstyrer modellen.", INK, BLUE, 10)
    merge(ws, "A5:K5", "Y/T-stykker og egne komponenter krever dokumentert K. K må referere til hastigheten i den valgte rørstrekningen. Ingen automatisk nettverksfordeling.", INK, AMBER, 10)
    headers = ["Type", "Vinkel [°]", "Antall", "Eget R/Di", "K, modell", "Eget K", "K, brukt", "Δp per enhet [kPa]", "Δp totalt [kPa]", "Kontroll", "Datakilde / kommentar"]
    for col, label in enumerate(headers, 1):
        styled_cell(ws, f"{chr(64+col)}6", label)
    ws.row_dimensions[6].height = 36
    defaults = {7: (FITTINGS[0][0],90,4), 8:(FITTINGS[0][0],45,2),
                9:(FITTINGS[4][0],90,0), 10:(FITTINGS[3][0],45,0), 11:(FITTINGS[1][0],90,0)}
    for r in range(FIRST,LAST+1):
        typ, angle, quantity = defaults.get(r,(None,None,0))
        for col, value in {"A":typ,"B":angle,"C":quantity,"D":None,"F":None,"K":None}.items():
            styled_cell(ws, f"{col}{r}", value, "input", "0" if col in ("B","C") else "0.000")
        radius = f'IF(D{r}="",IFERROR(VLOOKUP(A{r},Benddata,2,FALSE),0),D{r})'
        alpha, fd = f"RADIANS(B{r})", "Beregning!$B$23"
        sin = f"SIN({alpha}/2)"
        k = f'{fd}*{alpha}*({radius})+(0.1+2.4*{fd})*{sin}+6.6*{fd}*(SQRT({sin})+{sin})/(({radius})^(4*{alpha}/PI()))'
        styled_cell(ws, f"E{r}", f'=IFERROR(IF(AND(LEFT(A{r},4)="Bend",B{r}>0,B{r}<=180,({radius})>=0.5,Beregning!$B$3="OK"),{k},""),"")', number_format="0.0000")
        styled_cell(ws, f"G{r}", f'=IF(F{r}<>"",F{r},E{r})', number_format="0.0000")
        auto_status = '"OK"'
        for condition, message in reversed([
            (f'LEFT(A{r},4)<>"Bend"', "FEIL: oppgi eget K"),
            (f'OR(NOT(ISNUMBER(B{r})),B{r}<=0,B{r}>180,({radius})<0.5,NOT(ISNUMBER({radius})))', "FEIL: vinkel eller R/Di"),
            ('AND(Beregning!$B$14>0,Beregning!$B$14<4000)', "FEIL: eget K ved Re<4000"),
        ]):
            auto_status = f'IF({condition},"{message}",{auto_status})'
        own_status = f'IF(OR(NOT(ISNUMBER(F{r})),F{r}<0),"FEIL: eget K","OK")'
        status = f'IF(F{r}<>"",{own_status},{auto_status})'
        for condition, message in reversed([
            (f'OR(C{r}="",C{r}=0)', "Ikke med"),
            (f'OR(NOT(ISNUMBER(C{r})),C{r}<0,C{r}<>INT(C{r}))', "FEIL: antall"),
            (f'COUNTIF(Bendtyper,A{r})=0', "FEIL: velg type"),
            ('Beregning!$B$3<>"OK"', "Venter på inndata"),
        ]):
            status = f'IF({condition},"{message}",{status})'
        status = f'=IFERROR({status},"FEIL: kontroller tall")'
        styled_cell(ws, f"J{r}", status)
        styled_cell(ws, f"H{r}", f'=IF(J{r}="OK",G{r}*Beregning!$B$25/1000,"")', number_format="0.000")
        styled_cell(ws, f"I{r}", f'=IF(J{r}="OK",C{r}*H{r},"")', number_format="0.000")
    dv_list(ws, "Bendtyper", f"A{FIRST}:A{LAST}")
    dv_number(ws,f"B{FIRST}:B{LAST}",0,180)
    dv_number(ws,f"C{FIRST}:C{LAST}",0,100000,True)
    dv_number(ws,f"D{FIRST}:D{LAST}",0.5,1000)
    dv_number(ws,f"F{FIRST}:F{LAST}")
    total = LAST+2
    merge(ws, f"A{total}:H{total}", "SUM TRYKKTAP FRA BEND OG LOKALE MOTSTANDER", WHITE, TEAL, 11, True)
    styled_cell(ws, f"I{total}", f"=SUM(I{FIRST}:I{LAST})", number_format='0.000" kPa"')
    merge(ws, f"A{total+2}:K{total+3}", "Rennels-modellen er et estimat for rundede bend ved turbulent strømning. R/Di betyr bendets senterlinjeradius delt på faktisk innvendig diameter; en produsentverdi R/DN må omregnes. Modellen inkluderer friksjon i buen, så denne lengden skal ikke legges til rette rørmeter. Bruk leverandørens K ved annen geometri eller lave Reynolds-tall.", GREY, PALE, 10)
    ws.auto_filter.ref = f"A6:K{LAST}"
    ws.print_area = f"A1:K{total+3}"
    ws.conditional_formatting.add(f"J{FIRST}:J{LAST}", FormulaRule(formula=[f'LEFT(J{FIRST},4)="FEIL"'], fill=PatternFill("solid", fgColor=RED)))


def fixed_components(wb, name, title, descriptions):
    ws = base_sheet(wb,name,title,"Trykkfall skal være oppgitt ved volumstrømmen og driftsbetingelsene i hovedfanen.",[35,12,24,24,65,34])
    merge(ws,"A4:F4","Fyll inn antall og trykkfall PER ENHET i kPa. 1 bar = 100 kPa. Trykkfallene summeres som antall × Δp per enhet.",INK,BLUE,10)
    merge(ws,"A5:F5","Verdiene skaleres ikke automatisk når volumstrømmen endres. Oppdater dem fra leverandørdata. Antall 0 utelater raden.",INK,AMBER,10)
    headers=["Ventil / komponent","Antall","Δp per enhet [kPa]","Δp totalt [kPa]","Datakilde / kommentar","Kontroll"]
    for col,label in enumerate(headers,1):styled_cell(ws,f"{chr(64+col)}6",label)
    ws.row_dimensions[6].height=36
    for r in range(FIRST,LAST+1):
        styled_cell(ws,f"A{r}",descriptions[r-FIRST] if r-FIRST<len(descriptions) else None,"input")
        styled_cell(ws,f"B{r}",0,"input","0")
        styled_cell(ws,f"C{r}",None,"input","0.000")
        styled_cell(ws,f"E{r}",None,"input")
        status=f'=IFERROR(IF(OR(B{r}="",B{r}=0),"Ikke med",IF(OR(NOT(ISNUMBER(B{r})),B{r}<0,B{r}<>INT(B{r})),"FEIL: antall",IF(A{r}="","FEIL: angi navn",IF(OR(C{r}="",NOT(ISNUMBER(C{r})),C{r}<0),"FEIL: angi Δp per enhet","OK")))),"FEIL: kontroller tall")'
        styled_cell(ws,f"F{r}",status)
        styled_cell(ws,f"D{r}",f'=IF(F{r}="OK",B{r}*C{r},"")',number_format="0.000")
    dv_number(ws,f"B{FIRST}:B{LAST}",0,100000,True)
    dv_number(ws,f"C{FIRST}:C{LAST}")
    total=LAST+2
    merge(ws,f"A{total}:C{total}","SUM TRYKKTAP",WHITE,TEAL,11,True)
    styled_cell(ws,f"D{total}",f"=SUM(D{FIRST}:D{LAST})",number_format='0.000" kPa"')
    merge(ws,f"A{total+2}:F{total+3}","Registrer hver fysisk komponent én gang. En ventil som er lagt inn med trykkfall her, skal ikke i tillegg registreres som K-motstand i Bend-fanen.",GREY,PALE,10)
    ws.conditional_formatting.add(f"F{FIRST}:F{LAST}",FormulaRule(formula=[f'LEFT(F{FIRST},4)="FEIL"'],fill=PatternFill("solid",fgColor=RED)))
    ws.auto_filter.ref=f"A6:F{LAST}"
    ws.print_area=f"A1:F{total+3}"


def calculation(wb):
    ws=wb["Beregning"]
    for col,width in {"A":45,"B":31,"C":18,"D":35,"E":25}.items():ws.column_dimensions[col].width=width
    merge(ws,"A1:E2","BEREGNINGSGRUNNLAG",WHITE,NAVY,20,True)
    main="'Trykkberegning'!"
    def e(row):return f"{main}E{row}"
    checks=[
        (f'COUNTIF(Rorserier,{e(6)})=0',"FEIL: velg en gyldig rørserie."),
        (f'COUNTIF(Materialvalg,{e(13)})=0',"FEIL: velg materiale."),
        (f'COUNTIF(ENStandarder,{e(14)})=0',"FEIL: velg EN-produktstandard."),
        (f'AND({e(14)}<>"Annen / dokumentert",ISERROR(SEARCH("|"&{e(14)}&"|",VLOOKUP({e(13)},Materialdata,4,FALSE))))',"FEIL: EN-produktstandard passer ikke materialvalget. Se Materialer."),
        (f'COUNTIF(Medievalg,{e(20)})=0',"FEIL: velg medium."),
        (f'OR(NOT(ISNUMBER({e(21)})),{e(21)}<=0,NOT(ISNUMBER({e(22)})),{e(22)}<=0)',"FEIL: tetthet og dynamisk viskositet må være større enn 0."),
        (f'OR(NOT(ISNUMBER({e(12)})),{e(12)}<=0)',"FEIL: innvendig diameter må være større enn 0. Kontroller dimensjon og veggtykkelse."),
        (f'AND({e(11)}="",OR(NOT(ISNUMBER({e(10)})),{e(10)}<=0))',"FEIL: oppgi positiv veggtykkelse eller egen innvendig diameter."),
        (f'OR(NOT(ISNUMBER({e(8)})),{e(8)}<0,AND({e(8)}>0,{e(12)}>{e(8)}))',"FEIL: kontroller ytre og indre diameter."),
        (f'OR(NOT(ISNUMBER({e(15)})),{e(15)}<0,{e(15)}>={e(12)})',"FEIL: ruhet må være >=0 og mindre enn innvendig diameter."),
        (f'OR({e(17)}="",NOT(ISNUMBER({e(17)})),{e(17)}<0,{e(17)}<>INT({e(17)}))',"FEIL: flenskoblinger må være et heltall >=0."),
        (f'AND({e(17)}>0,OR({e(18)}="",NOT(ISNUMBER({e(18)})),{e(18)}<0))',"FEIL: oppgi K >=0 per flenskobling."),
        (f'OR({e(25)}="",NOT(ISNUMBER({e(25)})),{e(25)}<0)',"FEIL: volumstrøm må være et tall >=0."),
        (f'OR({e(29)}="",NOT(ISNUMBER({e(29)})),{e(29)}<0,{e(30)}="",NOT(ISNUMBER({e(30)})),{e(30)}<0,{e(31)}="",NOT(ISNUMBER({e(31)})),{e(31)}<0)',"FEIL: rette meter og høydemeter må være tall >=0."),
        (f'AND({e(33)}<>"",NOT(ISNUMBER({e(33)})))',"FEIL: innløpstrykk må være et tall eller tomt."),
    ]
    status='"OK"'
    for condition,message in reversed(checks):status=f'IF({condition},"{message}",{status})'
    base=f'=IFERROR({status},"FEIL: kontroller inndata; tallfelt kan ikke inneholde tekst.")'
    final='"OK"'
    row_count=LAST-FIRST+1
    final_checks=[('AND(B14>=2300,B14<4000)',"FEIL: Reynolds 2300–4000 er overgangsområde; modellen gir ikke et sikkert resultat.")]
    for sheet,col in [("Bend","J"),("Ventiler","F"),("Komponenter","F")]:
        cells=f"{sheet}!{col}{FIRST}:{col}{LAST}"
        final_checks.append((f'COUNTIF({cells},"OK")+COUNTIF({cells},"Ikke med")<>{row_count}',f"FEIL: kontroller aktive rader i {sheet}."))
    final_checks.append((f'AND(B7=0,SUM(Ventiler!D{FIRST}:D{LAST},Komponenter!D{FIRST}:D{LAST})>0)',"FEIL: ved null volumstrøm må komponenttrykkfall være 0."))
    for condition,message in reversed(final_checks):final=f'IF({condition},"{message}",{final})'
    final=f'=IFERROR(IF(B3<>"OK",B3,{final}),"FEIL: kontroller beregningsformler.")'
    entries={
        3:("Kontroll av grunninndata",base,""),4:("Samlet beregningsstatus",final,""),
        6:("Tyngdeakselerasjon g",9.80665,"m/s²"),
        7:("Volumstrøm Q",f"={e(25)}/3600","m³/s"),
        8:("Innvendig diameter D",f"={e(12)}/1000","m"),
        9:("Dynamisk viskositet μ",f"={e(22)}/1000","Pa·s"),
        10:("Tetthet ρ",f"={e(21)}","kg/m³"),
        11:("Absolutt ruhet ε",f"={e(15)}/1000","m"),
        12:("Tverrsnittsareal A","=PI()*B8^2/4","m²"),
        13:("Gjennomsnittshastighet v","=B7/B12","m/s"),
        14:("Reynolds-tall","=B10*B13*B8/B9","–"),
        15:("Relativ ruhet ε/D","=B11/B8","–"),
        16:("Startverdi f (Haaland / 64/Re)",'=IF(B14=0,0,IF(B14<2300,64/B14,(-1.8*LOG10((B15/3.7)^1.11+6.9/B14))^(-2)))',"–"),
        23:("Darcy friksjonsfaktor f","=B22","–"),
        24:("Strømningsregime",'=IF(B14=0,"Ingen strømning",IF(B14<2300,"Laminært",IF(B14<4000,"Overgang","Turbulent")))',""),
        25:("Dynamisk trykk ρv²/2","=B10*B13^2/2","Pa"),
        26:("Netto høydeforskjell",f"={e(30)}-{e(31)}","m"),
        29:("Rørfriksjon",f"=B23*{e(29)}/B8*B25","Pa"),
        30:("Bend og lokale K-motstander",f"=SUM(Bend!I{FIRST}:I{LAST})*1000","Pa"),
        31:("Ventiler",f"=SUM(Ventiler!D{FIRST}:D{LAST})*1000","Pa"),
        32:("Andre komponenter",f"=SUM(Komponenter!D{FIRST}:D{LAST})*1000","Pa"),
        33:("Flensede koblinger",f'=IF({e(17)}=0,0,{e(17)}*{e(18)}*B25)',"Pa"),
        35:("Dissipativt trykktap","=SUM(B29:B33)","Pa"),
        36:("Statisk høydebidrag","=B10*B6*B26","Pa"),
        37:("Nødvendig trykkdifferanse","=B35+B36","Pa"),
        38:("Colebrook-residual",'=IF(B14<2300,0,1/SQRT(B23)+2*LOG10(B15/3.7+2.51/(B14*SQRT(B23))))',"–"),
    }
    for row in range(17,23):
        entries[row]=(f"Colebrook-iterasjon {row-16}",f'=IF(B14=0,0,IF(B14<2300,B16,(-2*LOG10(B15/3.7+2.51/(B14*SQRT(B{row-1}))))^(-2)))',"–")
    for row,(label,value,unit) in entries.items():
        if isinstance(value,str) and value.startswith("=") and row not in (3,4):
            guard='B4="OK"' if row in (29,30,31,32,33,35,36,37) else 'B3="OK"'
            value=f'=IFERROR(IF({guard},{value[1:]},""),"")'
        styled_cell(ws,f"A{row}",label)
        styled_cell(ws,f"B{row}",value,number_format="0.000000")
        styled_cell(ws,f"C{row}",unit)
    for r in (3,4):ws.row_dimensions[r].height=85;ws[f"B{r}"].number_format="General"
    ws["B24"].number_format="General"
    styled_cell(ws,"D29","Bidrag")
    styled_cell(ws,"E29","bar")
    for row,label,source in [(30,"Rørfriksjon",29),(31,"Bend / K",30),(32,"Ventiler",31),(33,"Flenskoblinger",33),(34,"Andre komponenter",32),(35,"Høydebidrag",36)]:
        styled_cell(ws,f"D{row}",label)
        styled_cell(ws,f"E{row}",f'=IF(B4="OK",B{source}/100000,0)',number_format="0.000")
    ws.sheet_properties.tabColor=GREY
    ws.print_area="A1:E38"


def guide(wb):
    ws=base_sheet(wb,"Veiledning","SLIK BRUKER DU ARKET","Makrofri Excel-arbeidsbok · enheter og forutsetninger er en del av beregningen.",[8,42,115])
    topics=[
        ("1. Rør og materiale","Velg rørserie, DN/OD, materiale og EN-produktstandard. Oppgi faktisk veggtykkelse. Ved kjent lysåpning kan du overstyre ID. Metrisk OD er ikke en automatisk material- eller veggserie. En svart/grå verdi er beregnet; blå felt er inndata."),
        ("2. Væske og drift","Velg vann ved referansetemperatur eller annet medium. Overstyr tetthet og dynamisk viskositet ved faktisk drift. Viskositeten skal være i mPa·s (cP), ikke cSt. Volumstrøm angis i m³/h. Vann er standard; arket gjelder ikke gass eller damp."),
        ("3. Lengde og høyde","Rette rørmeter L er samlet rett rørlengde, både horisontal og vertikal. Høydefeltene legger ikke til rørfriksjonslengde. Oppgi positive høydemeter opp og ned. Netto høyde er opp minus ned. Fall reduserer nødvendig trykkdifferanse, men ikke dissipative tap. Arket løser ikke strømning fra tilgjengelig trykk; volumstrømmen er en inndata."),
        ("4. Bend / Y / T","Legg hver bendtype/vinkel på egen rad. R/Di er faktisk senterlinjeradius delt på innvendig diameter, ikke R/DN. Rundede bend bruker Rennels ved Re>=4000; eget K overstyrer. Y/T og annen geometri krever eget dokumentert K, fordi fordeling og strømningsvei påvirker motstanden. Maksimalt 25 rader."),
        ("Flensede rørkoblinger","I hovedfanen angis antall flensepar med pakning og K per kobling. K=0 forutsetter ideell full lysåpning uten innsnevring eller utstikkende pakning; flensen i seg selv gir da ikke et separat lokalt tap. Bruk dokumentert K ved avvik, eller registrer dokumentert trykkfall i Komponenter uten dobbelttelling. Flensens EN-standard/trykklasse, f.eks. EN 1092-1, inngår ikke i den hydrauliske beregningen."),
        ("5. Ventiler og komponenter","Velg antall og trykkfall per enhet i kPa i hver tabell. Trykkfall multipliseres med antall. Verdien skal gjelde den valgte volumstrømmen og åpningen/driftsbetingelsen. Den endres ikke automatisk med volumstrøm. Registrer ikke samme komponent også som K. Maksimalt 25 rader per tabell."),
        ("6. Resultater","Nødvendig trykkdifferanse = rørfriksjon + bendtap + ventiltap + øvrige tap + ρgΔz. Negativ verdi beholdes. Valgfritt innløpstrykk gir beregnet utløpstrykk. Beregningen forutsetter lik hastighet ved endene. Kontrollfeltet må vise OK; ved ugyldige aktive rader blir totalresultatet tomt."),
        ("7. Modellgrenser","Stasjonær, enfaset, Newtonsk og inkompressibel væske; én konstant diameter og volumstrøm; fullt utviklet rørstrømning. Re<2300 bruker 64/Re, Re>=4000 Colebrook. 2300<=Re<4000 blokkeres som overgangsområde. Rundede bends automatiske modell blokkeres ved 0<Re<4000; dokumentert eget K kan brukes. Ingen kavitasjons-, transient-, pumpekurve- eller nettverksberegning."),
        ("8. Normer og datakvalitet","EN/ISO-valgene er normreferanser og et begrenset dimensjonsutvalg, ikke samsvarsgodkjenning. Ruhet er et startanslag, ikke en materialgaranti. Arbeidsboken beregner hydraulisk trykktap; den beregner ikke tillatt arbeidstrykk eller veggtykkelse iht. EN 13480. Verifiser faktisk geometri og produktdata."),
        ("Darcy–Weisbach / Colebrook","Δp_rør = f(L/D)ρv²/2; Re=ρvD/μ; v=4Q/(πD²). Turbulent: 1/√f = −2 log10(ε/(3,7D)+2,51/(Re√f)). Seks iterasjoner fra Haaland. Laminært: f=64/Re. g=9,80665 m/s². Colebrook, C. F. (1939), Turbulent Flow in Pipes, with Particular Reference to the Transition Region between the Smooth and Rough Pipe Laws."),
        ("Rennels, rundet bend","α i radianer, r=R/Di: K=fαr+(0,10+2,4f)sin(α/2)+6,6f[√sin(α/2)+sin(α/2)]/r^(4α/π). Modellen inkluderer friksjon i den buede rørlengden. Rennels & Hudson, Pipe Flow: A Practical and Comprehensive Guide, Wiley (2012). Koden er kontrollert mot den offentlig dokumenterte implementasjonen i fluids.fittings.bend_rounded."),
        ("Referanser / etterprøving","Normreferanser: EN 10220, ISO 4200, ISO 1127; relevante deler av EN 10216 og EN 10217; EN 10088-1 for rustfrie materialbetegnelser. Vann: avrundede standard tabellverdier (IAPWS); ISO VG 32: nominell viskositetsklasse iht. ISO 3448. Oppslagsarkene viser alle verdier brukt; testskriptet kontrollerer også uavhengig Colebrook-løsning og laminær Poiseuille."),
    ]
    for row,(topic,text) in enumerate(topics,6):
        styled_cell(ws,f"A{row}",row-5)
        styled_cell(ws,f"B{row}",topic)
        styled_cell(ws,f"C{row}",text)
        ws.row_dimensions[row].height=90 if row in (12,14,15,16) else 78
    row=19
    cell=styled_cell(ws,f"C{row}","Dokumentert referanseimplementasjon: fluids / fittings.py")
    cell.hyperlink="https://github.com/CalebBell/fluids/blob/master/fluids/fittings.py"
    cell.font=Font(color="086AA5",underline="single")
    ws.print_area="A1:C19"
    ws.page_setup.fitToHeight=0
    ws.print_title_rows="1:3"


def make_workbook():
    wb=Workbook()
    wb.remove(wb.active)
    wb.properties.title="Trykkberegning av rør – væsker"
    wb.properties.subject="Darcy–Weisbach, Colebrook, bend og komponenttap"
    wb.properties.creator="R-rdimensjonering"
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    wb.create_sheet("Beregning")
    references(wb)
    inputs(wb)
    fitting_rows(wb)
    fixed_components(wb,"Ventiler","VENTILER – ANTALL OG TRYKKFALL",["Kuleventil","Reguleringsventil","Tilbakeslagsventil","Spjeldventil","Sluseventil"])
    fixed_components(wb,"Komponenter","ANDRE KOMPONENTER – TRYKKFALL",["Filter","Varmeveksler","Mengdemåler","Sil","Annen komponent"])
    calculation(wb)
    guide(wb)
    order=["Trykkberegning","Bend","Ventiler","Komponenter","Veiledning","Rørdata","Materialer","Medier","Beregning","_Lister"]
    for i,name in enumerate(order):wb.move_sheet(wb[name],offset=i-wb.index(wb[name]))
    wb.active=0
    for sheet in wb:
        sheet.sheet_view.showGridLines=False
    return wb


def restore_validations(source, target):
    """Preserve Excel validation expressions without rewriting cached formulas.

    Calc drops some INDIRECT-based validation lists on XLSX export. Restore only
    the validation XML from our generated source; all calculated cell XML stays
    byte-for-byte intact outside that block.
    """
    main_ns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ET.register_namespace("",main_ns)
    def sheet_paths(package):
        book=ET.fromstring(package.read("xl/workbook.xml"))
        rels=ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        targets={r.attrib["Id"]:r.attrib["Target"] for r in rels}
        result={}
        for sheet in book.find(f"{{{main_ns}}}sheets"):
            relative=targets[sheet.attrib[f"{{{rel_ns}}}id"]]
            result[sheet.attrib["name"]]=relative.lstrip("/") if relative.startswith("/") else str(PurePosixPath("xl")/relative)
        return result
    changes={}
    with ZipFile(source) as original,ZipFile(target) as calculated:
        originals,outputs=sheet_paths(original),sheet_paths(calculated)
        for name,path in originals.items():
            tree=ET.fromstring(original.read(path))
            validation=tree.find(f"{{{main_ns}}}dataValidations")
            if validation is None:continue
            block=ET.tostring(validation,encoding="utf-8")
            output_path=outputs[name]
            content=calculated.read(output_path)
            pattern=rb"<dataValidations\b[^>]*>.*?</dataValidations>"
            content,count=re.subn(pattern,lambda _:block,content,flags=re.DOTALL)
            if count==0:
                # These worksheet elements follow dataValidations in OOXML.
                match=re.search(rb"<(?:hyperlinks|printOptions|pageMargins|pageSetup|headerFooter|drawing|legacyDrawing|tableParts|extLst)\b",content)
                position=match.start() if match else content.index(b"</worksheet>")
                content=content[:position]+block+content[position:]
            elif count!=1:raise RuntimeError("Uventet antall dataValidations-blokker")
            changes[output_path]=content
        with tempfile.NamedTemporaryFile(dir=target.parent,suffix=".xlsx",delete=False) as temporary:
            temporary_path=Path(temporary.name)
        try:
            with ZipFile(temporary_path,"w") as rebuilt:
                rebuilt.comment=calculated.comment
                for item in calculated.infolist():rebuilt.writestr(item,changes.get(item.filename,calculated.read(item.filename)))
        except BaseException:
            temporary_path.unlink(missing_ok=True)
            raise
    temporary_path.replace(target)


def recalculate(inputs, output_dir, profile, timeout=240):
    if not inputs:return
    executable=shutil.which("soffice")
    if not executable:raise RuntimeError("LibreOffice/soffice mangler: kreves for å beregne lagrede formelresultater.")
    output_dir.mkdir(parents=True,exist_ok=True)
    cache=profile.parent/"cache"
    cache.mkdir(parents=True,exist_ok=True)
    runtime_env=os.environ.copy()
    runtime_env["XDG_CACHE_HOME"]=str(cache.resolve())
    command=[executable,f"-env:UserInstallation={profile.resolve().as_uri()}","--headless","--convert-to","xlsx:Calc MS Excel 2007 XML","--outdir",str(output_dir),*[str(p) for p in inputs]]
    # soffice launches a native child; stop the whole owned group on timeout.
    with subprocess.Popen(command,env=runtime_env,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,start_new_session=True) as process:
        try:
            stdout,stderr=process.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            os.killpg(process.pid,signal.SIGTERM)
            try:process.communicate(timeout=10)
            except subprocess.TimeoutExpired:
                os.killpg(process.pid,signal.SIGKILL)
                process.communicate()
            raise
        if process.returncode:
            raise subprocess.CalledProcessError(process.returncode,command,stdout,stderr)
    for path in inputs:
        target=output_dir/path.name
        if not target.is_file():raise RuntimeError(f"Beregnet arbeidsbok mangler: {path.name}\n{stdout}\n{stderr}")
        restore_validations(path,target)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output",type=Path,default=ROOT/"workbooks"/"Trykkberegning_ror.xlsx")
    parser.add_argument("--overwrite",action="store_true",help="Erstatt eksisterende generert fil; ikke bruk på utfylte brukerfiler.")
    args=parser.parse_args()
    if args.output.exists() and not args.overwrite:parser.error("Filen finnes allerede. Velg ny --output eller bruk --overwrite for å regenerere malen.")
    with tempfile.TemporaryDirectory(prefix="pipe-workbook-") as directory:
        root=Path(directory)
        source=root/"input";source.mkdir()
        raw=source/args.output.name
        wb=make_workbook();wb.save(raw);wb.close()
        recalculate([raw],root/"calculated",root/"profile")
        args.output.parent.mkdir(parents=True,exist_ok=True)
        shutil.copyfile(root/"calculated"/raw.name,args.output)
    print(f"Lagret beregnet Excel-mal: {args.output}")


if __name__=="__main__":main()
