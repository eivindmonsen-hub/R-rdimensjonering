"""Validate Excel file handling and formula recalculation outside the repository."""
import math
import shutil
import subprocess
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def main():
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("soffice is required for formula recalculation")

    with tempfile.TemporaryDirectory(prefix="excel-readiness-") as directory:
        root = Path(directory)
        source = root / "input"
        target = root / "recalculated"
        source.mkdir()
        target.mkdir()
        path = source / "regneark.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Kontroll"
        sheet.append(["Beskrivelse", "Verdi"])
        sheet.append(["Vannføring (m³/h)", 36])
        sheet.append(["Diameter (m)", 0.1])
        sheet.append(["Vannføring (m³/s)", "=B2/3600"])
        sheet.append(["Areal (m²)", "=PI()*B3^2/4"])
        sheet.append(["Hastighet (m/s)", "=B4/B5"])
        sheet.append(["Kontrollsum", "=SUM(B2:B3)"])
        sheet.append(["Formelstatus", '=IF(B6>0,"OK","FEIL")'])
        sheet["A1"].font = Font(bold=True)
        sheet["B6"].number_format = "0.0000"
        workbook.save(path)
        workbook.close()

        original = load_workbook(path)
        assert original["Kontroll"]["B6"].value == "=B4/B5"
        assert original["Kontroll"]["A1"].font.bold
        original.close()

        result = subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={(root / 'profile').as_uri()}",
                "--headless",
                "--convert-to",
                "xlsx:Calc MS Excel 2007 XML",
                "--outdir",
                str(target),
                str(path),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=90,
        )
        output = target / path.name
        if not output.is_file():
            raise RuntimeError(f"No recalculated workbook: {result.stdout} {result.stderr}")

        calculated = load_workbook(output, data_only=True)
        values = calculated["Kontroll"]
        expected = {"B4": 0.01, "B5": math.pi * 0.1**2 / 4, "B6": 4 / math.pi, "B7": 36.1}
        for cell, expected_value in expected.items():
            actual = values[cell].value
            assert isinstance(actual, (int, float)), (cell, actual)
            assert math.isclose(actual, expected_value, rel_tol=1e-10), (cell, actual, expected_value)
        assert values["B8"].value == "OK"
        assert values["A2"].value == "Vannføring (m³/h)"
        calculated.close()

        formulas = load_workbook(output, data_only=False)
        assert formulas["Kontroll"]["B6"].data_type == "f"
        formulas.close()
        frame = pd.read_excel(output, sheet_name="Kontroll", engine="openpyxl")
        assert len(frame) == 7
        assert frame.iloc[6]["Verdi"] == "OK"
        print("PASS: XLSX writing, reading, formatting, Unicode, four numerical formulas, IF, formula preservation and pandas import")


if __name__ == "__main__":
    main()
